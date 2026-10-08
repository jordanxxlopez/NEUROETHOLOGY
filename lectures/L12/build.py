#!/usr/bin/env python3
"""Run the repository builder, complete non-content notes, and recheck."""
import json,subprocess,sys
from pathlib import Path
from pptx import Presentation
BASE=Path(__file__).resolve().parent
ROOT=BASE.parents[1]
subprocess.run([sys.executable,str(BASE/'write_spec.py')],check=True)
subprocess.run([sys.executable,str(ROOT/'tools/build_lecture.py'),str(BASE/'lecture.json')],check=True)
spec=json.loads((BASE/'lecture.json').read_text())
p=BASE/'Neuroethology_Lecture12_FA2026.pptx'
prs=Presentation(p)
prs.slides[0].notes_slide.notes_text_frame.text+='\n\nOriginal article photograph credit: Colin Hutton, credited by Zurek et al. (2015) for Fig. 1(A). The photograph retains its published pixels.'
tk=spec['takeaways']
prs.slides[-1].notes_slide.notes_text_frame.text='\n'.join('• '+x for x in tk['transcript'])+'\n\nReferences:\n'+'\n'.join(tk['refs'])
for slide in prs.slides:
 for para in slide.notes_slide.notes_text_frame.paragraphs:
  for run in para.runs:run.font.name='Arial'
prs.save(p)
subprocess.run([sys.executable,str(ROOT/'tools/check_lecture.py'),str(p),'--lecture','12'],check=True)
