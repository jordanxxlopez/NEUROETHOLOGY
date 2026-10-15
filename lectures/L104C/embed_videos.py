"""Place original figures, licensed companion photographs and native research movies."""
from pathlib import Path
import json,hashlib
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.dml.color import RGBColor
R=Path(__file__).resolve().parent;P=R/'Neuroethology_Lecture104C_FA2026.pptx';spec=json.load(open(R/'lecture.json'));prs=Presentation(P);rows=json.load(open(R/'media_preparation.json'));refs=json.load(open(R/'reference_verification.json'))
assert len(prs.slides)==51
assert not any(sh.shape_type==MSO_SHAPE_TYPE.MEDIA for sl in prs.slides for sh in sl.shapes),'Rebuild before embedding.'
mapping={'sample_video_ant_140_visit_2':(12,'Dataset sample_video_ant_140_visit_2.mp4','Ant 140, visit 2, 250 ppm caffeine; original recording excerpt 10–100 s.'),'sample_video_ant_140_visit_5':(13,'Dataset sample_video_ant_140_visit_5.mp4','Ant 140, visit 5, 250 ppm caffeine; original recording excerpt 10–100 s.'),'sample_video_ant_40_visit_4':(14,'Dataset sample_video_ant_40_visit_4.mp4','Ant 40, visit 4, 2000 ppm caffeine; original recording excerpt 10–100 s.'),'sample_video_ant_40_visit_5':(15,'Dataset sample_video_ant_40_visit_5.mp4','Ant 40, visit 5, 2000 ppm caffeine; original recording excerpt 10–100 s.'),'WormAlcohol_Movie1':(21,'S1 Movie','Wild type after ethanol exposure.'),'WormAlcohol_Movie2':(22,'S2 Movie','dop-2 mutant movement under ethanol.'),'WormAlcohol_Movie6':(23,'S6 Movie','dop-2 movement after PDE ablation under ethanol.'),'WormAlcohol_Movie7':(24,'S7 Movie','Wild type with dopamine, without ethanol.'),'WormAlcohol_Movie8':(24,'S8 Movie','Wild type with dopamine and ethanol.'),'WormAlcohol_Movie9':(25,'S9 Movie','NLP-12 overexpression under ethanol.'),'FlyDopamine_Movie1':(26,'Movie S1','Confocal reconstruction of labeled PPM3 projections.'),'FlyCourtship_Movie1':(28,'Movie S1','Wild-type courtship before and across ethanol exposures.'),'FlyCourtship_Movie2':(29,'Movie S2','Courtship under ethanol in the compared white genotypes.'),'FishNicotine_Movie1':(34,'Movie S1','Larval locomotion in the nicotine-response assay.'),'FishNicotine_Movie2':(35,'Movie S2','Article-provided larval locomotion recording.'),'FishNicotine_Movie3':(36,'Movie S3','Article-provided larval locomotion recording.'),'FishNicotine_Movie4':(37,'Movie S4','Three-dimensional neural reporter fluorescence.'),'MouseEthanol_Movie1':(43,'Movie S1','Baseline mitochondrial transport in cortical axons.'),'MouseEthanol_Movie2':(44,'Movie S2','Mitochondrial transport following ethanol intoxication.'),'MouseEthanol_Movie3':(45,'Movie S3','Dense-core-vesicle transport in cortical axons.')}
by_slide={};sources=[]
for row in rows:
 stem=Path(row['delivery_file']).stem;n,ident,desc=mapping[stem];ref=refs[row['article_doi'].lower()];author=ref['authorString'].split()[0];author='Galante' if stem.startswith('sample_') else author;caption=f"{author} et al. ({ref['pubYear']}), {ident}. {desc}";row=dict(row,slide=n,original_identifier=ident,description=desc,caption=caption,credit=ref['authorString'],embedded=True);by_slide.setdefault(n,[]).append(row);sources.append(row)
def fit(pic,x,y,w,h):
 ratio=pic.width/pic.height;fw=min(w,h*ratio);fh=fw/ratio;pic.left=Inches(x+(w-fw)/2);pic.top=Inches(y+(h-fh)/2);pic.width=Inches(fw);pic.height=Inches(fh)
def caption(sl,text,x,y,w,h,size=9,name='Figure caption'):
 b=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));b.name=name;tf=b.text_frame;tf.word_wrap=True;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.text=text
 for p in tf.paragraphs:
  p.space_after=Pt(0)
  for ru in p.runs:ru.font.name='Arial';ru.font.size=Pt(size);ru.font.color.rgb=RGBColor(0,0,0)
 return b
photo_placements=[]
for n,s in enumerate(spec['slides'],2):
 sl=prs.slides[n-1];teaching=[x.text for x in sl.shapes if x.has_text_frame and x.name!='Figure caption'];pics=[x for x in sl.shapes if x.shape_type==MSO_SHAPE_TYPE.PICTURE];oldcap=[x for x in sl.shapes if x.name=='Figure caption'];figs=[s['figure']] if 'figure' in s else s['figures'];x=13.333-.62-s['figure_width'];w=s['figure_width'];video=n in by_slide
 # Existing captions are recreated in defined slots; no scientific pixels are edited.
 for sh in oldcap:sh._element.getparent().remove(sh._element)
 companion=s.get('companion');slots=len(pics)+(1 if companion else 0)
 if video:
  top_h=2.45;cap_y=3.88;cap_h=.52
  if slots==1:locs=[(x,w)]
  else:locs=[(x,w-1.7),(x+w-1.5,1.5)]
  for j,pic in enumerate(pics):
   xx,ww=locs[j];fit(pic,xx,1.36,ww,top_h);caption(sl,figs[j]['caption'],xx,cap_y,ww,cap_h,8)
 else:
  top_h=4.70;cap_y=6.15;cap_h=.65
  if slots==1:locs=[(x,w)]
  elif len(pics)==2 and not companion:
   # Evidence and its anatomical locator occupy separate vertical regions.
   for j,pic in enumerate(pics):
    yy=1.36+j*2.66;fit(pic,x,yy,w,2.04);caption(sl,figs[j]['caption'],x,yy+2.08,w,.48,8.5)
   locs=[]
  else:locs=[(x,w-1.7),(x+w-1.5,1.5)]
  if locs:
   for j,pic in enumerate(pics):
    xx,ww=locs[j];fit(pic,xx,1.36,ww,top_h);caption(sl,figs[j]['caption'],xx,cap_y,ww,cap_h,9)
 if companion:
  xx,ww=locs[-1];photo=sl.shapes.add_picture(str(R/companion['path']),Inches(xx),Inches(1.36));fit(photo,xx,1.36,ww,top_h);photo.name='Additional credited subject image'
  # A concise on-slide credit accompanies the complete license and source in notes.
  credit=companion['credit'].replace('(c) ','').split(',')[0];txt=companion['caption']+' '+credit+'. '+companion['license']+'.';caption(sl,txt,xx,cap_y,ww,cap_h,8)
  sl.notes_slide.notes_text_frame.text+='\n\nImage: '+companion['caption']+' Credit: '+companion['credit']+'. License: '+companion['license']+'. Source: '+companion['source_url'];photo_placements.append(dict(slide=n,**companion))
 for j,row in enumerate(by_slide.get(n,[])):
  file=R/row['delivery_file'];assert hashlib.sha256(file.read_bytes()).hexdigest()==row['delivery_sha256'];ratio=row['width']/row['height'];count=len(by_slide[n]);slotw=(w-.18*(count-1))/count;maxh=1.98;mw=min(slotw,maxh*ratio);mh=mw/ratio;mx=x+j*(slotw+.18)+(slotw-mw)/2;my=4.49+(maxh-mh)/2
  fallback=sl.shapes.add_picture(str(R/row['poster']),Inches(mx),Inches(my),width=Inches(mw),height=Inches(mh));fallback.name='Static video fallback'
  m=sl.shapes.add_movie(str(file),Inches(mx),Inches(my),Inches(mw),Inches(mh),poster_frame_image=str(R/row['poster']),mime_type='video/mp4');m.name=Path(row['delivery_file']).stem
  for el in m._element.xpath('.//p:cNvPr'):el.set('descr',row['caption']+' Click to play the embedded recording.')
  for el in sl._element.xpath('.//p:video'):el.set('fullScrn','1')
  caption(sl,row['caption'],x+j*(slotw+.18),6.55,slotw,.39,8.3,'Video caption');row['native_aspect_ratio']=ratio
  url=row.get('url',row.get('download_url',row.get('original_url','')))
  sl.notes_slide.notes_text_frame.text+='\n\nEmbedded video: '+row['caption']+'\nExact identifier: '+row['original_identifier']+'\nArticle DOI: https://doi.org/'+row['article_doi']+'\nOriginal URL: '+url+'\nCredit: '+row['credit']+'\nLicense: '+row.get('license','Original article recording; no additional open license stated')+'\nConversion: '+row.get('conversion','Compatibility conversion; original scientific content preserved.')+'\nPlayback: native embedded MP4, fullscreen flag enabled. Desktop PowerPoint GUI playback and Google Slides import have not been tested. Google Slides may require separate insertion of the video from Drive.\n'
 assert [x.text for x in sl.shapes if x.has_text_frame and x.name not in ('Figure caption','Video caption')]==teaching
# Complete title-photograph credit, including its license.
t=spec['title_image'];prs.slides[0].notes_slide.notes_text_frame.text+='\n\nImage: '+t['caption']+' Credit: '+t['credit']+'. License: '+t['license']+'. Source: '+t['source_url']
# Rendered QA: preserve font sizes and prose while removing three small paragraph-spacing overruns.
for n in (19,22,28):
 body=next(sh for sh in prs.slides[n-1].shapes if sh.name=='Body')
 for para in body.text_frame.paragraphs:para.space_after=Pt(5)
prs.save(P)
(R/'media_sources.json').write_text(json.dumps(dict(embedded_video_count=len(sources),distinct_original_recordings=len({r.get('original_sha256',r.get('sha256')) for r in sources}),sources=sources),indent=2,ensure_ascii=False)+'\n');(R/'additional_image_placements.json').write_text(json.dumps(photo_placements,indent=2,ensure_ascii=False)+'\n');print('Embedded',len(sources),'original recordings; added',len(photo_placements),'credited photographs; PPTX',round(P.stat().st_size/1048576,1),'MiB')
