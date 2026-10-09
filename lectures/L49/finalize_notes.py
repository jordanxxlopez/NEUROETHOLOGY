"""Apply Arial and plain, readable teaching text to the repository-built notes."""
from pathlib import Path
import sys
from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor
ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT/'tools'))
from build_lecture import plain
path=Path(__file__).with_name('Neuroethology_Lecture49_FA2026.pptx')
prs=Presentation(path)
# Fit the exact long special-topic title and instructor within the title panel.
for shape in prs.slides[0].shapes:
 if not shape.has_text_frame:continue
 if shape.name == 'Lecture title':shape.top=Inches(2.25)
 elif shape.text == 'Lecture 49':shape.top=Inches(1.65)
 elif shape.text == 'Date TBD':shape.top=Inches(5.9)
 elif shape.text.startswith('Jordan Lopez'):shape.top=Inches(6.65)
for index, slide in enumerate(prs.slides):
 if 0 < index < len(prs.slides)-1:
  body = next(s for s in slide.shapes if s.has_text_frame and len(s.text) > 600)
  for paragraph in body.text_frame.paragraphs:
   for run in paragraph.runs:
    if run.font.size and run.font.size.pt > 14.5:
     run.font.size = Pt(14.5)
 for paragraph in slide.notes_slide.notes_text_frame.paragraphs:
  paragraph.text=plain(paragraph.text)
  for run in paragraph.runs:
   run.font.name='Arial'
   run.font.size=Pt(12)
   run.font.color.rgb=RGBColor(0,0,0)
prs.save(path)
print('Applied Arial to all speaker notes and removed literal Markdown emphasis markers.')
