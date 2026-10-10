"""Build through tools/build_lecture.py with spacing for the exact long title."""
import importlib.util,json,sys
from pathlib import Path
from pptx.util import Inches
D=Path(__file__).resolve().parent;R=D.parents[1];sys.path.insert(0,str(R/'tools'));sp=importlib.util.spec_from_file_location('lecture_builder',R/'tools/build_lecture.py');b=importlib.util.module_from_spec(sp);sp.loader.exec_module(b)
original=b.Deck.title_slide
def title_slide(self):
 original(self);s=self.prs.slides[0]
 for shape in s.shapes:
  if not shape.has_text_frame:continue
  text=shape.text
  if text==self.course['course']:shape.top=Inches(.45)
  elif text==f"Lecture {self.meta['n']}":shape.top=Inches(1.05)
  elif shape.name=='Lecture title':shape.top=Inches(1.6)
  elif text==self.meta['date']:shape.top=Inches(6.4)
  elif text.startswith(self.course['instructor']):shape.top=Inches(6.85);shape.height=Inches(.6)
b.Deck.title_slide=title_slide
sys.argv=['build_lecture.py',str(D/'lecture.json')];b.main()
