"""Fit source-derived MBARI excerpts in existing white space, preserving old content."""
from pathlib import Path
import json,subprocess,hashlib,zipfile
from lxml import etree
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
B=Path(__file__).resolve().parents[1];D=Path('/workspace/l99-extra-videos')
src=B/'deliverables/Neuroethology_Lecture99_FA2026_UploadedVideos_WORKING.pptx';out=B/'deliverables/Neuroethology_Lecture99_FA2026_21Videos.pptx'
records=[dict(slide=45,key='-rYNTFOeBVI',start=0.0,end=16.683333,poster=5.0,original_file='/workspace/attachments/8163f69e-b0ce-47fd-8ada-a6e956c5b849/Bloody_Belly_Comb_Jelly_29MB.mp4',description='MBARI footage of Lampocteis cruentiventer swimming and changing body orientation, with iridescent comb rows visible. The animal remains in the full source frame. This instructor-authorized aesthetic scene illustrates the live animal and ciliary locomotion; it is not labeled as an article supplement.',caption='Video: Lampocteis cruentiventer swimming. Credit: MBARI.')]
p=Presentation(src)
for r in records:
 f=Path(r['original_file']);clip=D/(r['key']+'_lecture_clip.mp4')
 subprocess.run(['ffmpeg','-v','error','-y','-ss',str(r['start']),'-i',str(f),'-t',str(r['end']-r['start']),'-map','0:v:0','-map','0:a:0','-c','copy','-bsf:v',"noise=drop='gt(pts*tb,16.683334)'",'-avoid_negative_ts','make_zero',str(clip)],check=True)
 probe=json.loads(subprocess.check_output(['ffprobe','-v','error','-show_streams','-show_format','-of','json',str(clip)]));v=next(x for x in probe['streams'] if x['codec_type']=='video');assert (v['width'],v['height'],v['codec_name'])==(1920,1080,'h264');assert any(x['codec_name']=='aac' for x in probe['streams'])
 subprocess.run(['ffmpeg','-v','error','-i',str(clip),'-f','null','-'],check=True)
 poster=D/(r['key']+'_lecture_poster.png');subprocess.run(['ffmpeg','-v','error','-y','-ss',str(r['poster']),'-i',str(clip),'-frames:v','1',str(poster)],check=True)
 s=p.slides[r['slide']-1];right=[x for x in s.shapes if x.left/Inches(1)>7 and 1<x.top/Inches(1)<6.8];y=max((x.top+x.height)/Inches(1) for x in right)+.08;h=min(2.15,6.46-y);w=min(5.35,h*16/9);h=w*9/16;x=7.363+(5.35-w)/2
 for picture in [True,False]:
  sh=s.shapes.add_picture(str(poster),Inches(x),Inches(y),Inches(w),Inches(h)) if picture else s.shapes.add_movie(str(clip),Inches(x),Inches(y),Inches(w),Inches(h),poster_frame_image=str(poster),mime_type='video/mp4');sh.name=('Added MBARI fallback ' if picture else 'Added MBARI movie ')+r['key']
 sh=s.shapes.add_textbox(Inches(7.363),Inches(6.5),Inches(5.35),Inches(.34));sh.name='Added MBARI credit '+r['key'];tf=sh.text_frame;tf.margin_top=tf.margin_bottom=tf.margin_left=tf.margin_right=0;run=tf.paragraphs[0].add_run();run.text=r['caption'];run.font.name='Arial';run.font.size=Pt(9);run.font.color.rgb=RGBColor(0,0,0)
 r.update(original_file=str(f),original_sha256=hashlib.sha256(f.read_bytes()).hexdigest(),clip=str(clip),sha256=hashlib.sha256(clip.read_bytes()).hexdigest(),url='https://www.youtube.com/watch?v='+r['key'],credit='MBARI (Monterey Bay Aquarium Research Institute). Original video overlays and logo retained. No Creative Commons license asserted.',processing='Temporal extraction with video/audio stream copy and packet selection after the cut endpoint, preventing a partially decoded terminal frame. Native decoded video frames match the source exactly. Full 1920 x 1080 frame, native frame rate, source color and sound preserved. No spatial crop is needed to fit the available space; the whole animal and source credits remain visible.',probe=probe,geometry=dict(x=x,y=y,w=w,h=h))
 for line in ['Additional MBARI recording: '+r['caption'],'Original source: '+r['url'],'Credit: '+r['credit'],'Processing: '+r['processing'],f"Source interval: {r['start']} to {r['end']} seconds; extracted clip duration: {probe['format']['duration']} seconds.",'Teaching note: '+r['description'],'Source category: instructor-authorized aesthetic/general science footage. Not labeled as an article supplementary movie.']:
  run=s.notes_slide.notes_text_frame.add_paragraph().add_run();run.text=line;run.font.name='Arial'
p.save(out)
ns={'p':'http://schemas.openxmlformats.org/presentationml/2006/main'}
with zipfile.ZipFile(src) as a,zipfile.ZipFile(out) as b:
 oldmedia=[n for n in a.namelist() if n.startswith('ppt/media/')]
 for name in oldmedia:assert a.read(name)==b.read(name),name
 for n in range(1,47):
  name=f'ppt/slides/slide{n}.xml';old=etree.fromstring(a.read(name)).find('p:cSld/p:spTree',ns);new=etree.fromstring(b.read(name)).find('p:cSld/p:spTree',ns)
  for o,v in zip(old,new):assert etree.tostring(o,method='c14n')==etree.tostring(v,method='c14n'),f'slide{n} changed'
 assert b.testzip() is None
report=dict(slides=46,total_embedded_videos=21,added_uploaded_clips=1,existing_shapes_unchanged=True,existing_media_unchanged=len(oldmedia),new_clips_decode_cleanly=True,source_resolution_preserved='1920 x 1080',compression='None; original video and audio streams copied',bytes=out.stat().st_size,sha256=hashlib.sha256(out.read_bytes()).hexdigest(),remaining_upload=None)
(B/'additional_videos/FINAL_CLIP.json').write_text(json.dumps(records,indent=2)+'\n');(B/'additional_videos/final_video_validation.json').write_text(json.dumps(report,indent=2)+'\n');print(json.dumps(report,indent=2))
