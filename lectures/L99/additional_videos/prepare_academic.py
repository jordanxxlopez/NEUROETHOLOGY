"""Add seven academic recordings without changing any existing slide shape or media."""
from pathlib import Path
import json,subprocess,hashlib,zipfile
from lxml import etree
from pptx import Presentation
from pptx.util import Inches,Pt
D=Path('/workspace/l99-extra-videos');B=Path(__file__).resolve().parents[1]
src=B/'deliverables/Neuroethology_Lecture99_FA2026.pptx';out=B/'deliverables/Neuroethology_Lecture99_FA2026_AdditionalVideos_WORKING.pptx'
records=[]
for i,n,desc in [(4,32,'Peptide-labeled cells in the aboral organ, in three dimensions.'),(5,34,'Peptide-labeled pharyngeal and aboral tissue.'),(6,28,'Three-dimensional localization of peptide-labeled cells.'),(7,33,'Peptide-labeled aboral tissue viewed from multiple angles.'),(8,25,'Electron-microscopy reconstruction of the surface nerve net and nearby cells.')]:
 records.append(dict(slide=n,path=str(D/f'Sachkova_mmc{i}.mp4'),doi='10.1016/j.cub.2021.09.005',url=f'https://ars.els-cdn.com/content/image/1-s2.0-S0960982221012458-mmc{i}.mp4',caption=f'Sachkova (2021), Video S{i-3}. '+desc,description=desc,credit='Sachkova et al.; original Current Biology supplementary recording. Publisher terms apply.'))
records.extend([dict(slide=9,path=str(D/'download0.mp4'),doi='10.1016/j.cub.2022.09.061',url='https://ars.els-cdn.com/content/image/1-s2.0-S0960982222016037-mmc2.mp4',caption='Jokura (2022), Video S1. Control and CTENO189 knockdown under viscous load.',description='Control versus CTENO189 morpholino recordings, with original condition labels and scale bars intact.',credit='Jokura et al.; original Current Biology supplementary recording. Publisher terms apply.'),dict(slide=4,path=str(D/'Jokura2019_VideoS1.mp4'),doi='10.1016/j.cub.2019.08.059',url='https://ars.els-cdn.com/content/image/1-s2.0-S0960982219311091-mmc2.mp4',caption='Jokura (2019), article-provided video: mmc2.mp4. Comb rows of a swimming ctenophore.',description='Whole-animal comb-row beating, showing the locomotor plates in action. The 2022 methods explicitly link this earlier study to an author-provided YouTube recording.',credit='Jokura et al.; original Current Biology supplementary recording. Publisher terms apply.')])
ref=json.loads((D/'more0.dat').read_text())['message']['items'][1]
full_reference=', '.join(a.get('family','') for a in ref['author'])+' (2019). '+ref['title'][0]+'. '+ref['container-title'][0]+' '+ref.get('volume','')+', '+ref.get('page','')+'. https://doi.org/'+ref['DOI']
p=Presentation(src)
for r in records:
 f=Path(r['path']);pr=json.loads(subprocess.check_output(['ffprobe','-v','error','-show_streams','-show_format','-of','json',str(f)]));vs=next(v for v in pr['streams'] if v['codec_type']=='video' and not v.get('disposition',{}).get('attached_pic'))
 # Keep source-authored visualizations exactly as published. Trim only the long EM recording.
 embed=f;r['processing']='Original file embedded byte-for-byte; no resizing, recompression, recoloring or speed change.'
 if r['slide']==25:
  embed=D/'Sachkova_S5_first30s.mp4';subprocess.run(['ffmpeg','-v','error','-y','-i',str(f),'-t','30','-map','0:v:0','-map','0:a:0?','-c','copy',str(embed)],check=True);r['processing']='First 30 seconds, video and audio stream copied without re-encoding; no resizing or speed change.'
 poster=D/f"poster_slide{r['slide']}.png";subprocess.run(['ffmpeg','-v','error','-y','-ss',('2.0' if r['slide']==9 else '0.5'),'-i',str(embed),'-map','0:v:0' if r['slide']!=4 else '0:v:0','-frames:v','1',str(poster)],check=True)
 s=p.slides[r['slide']-1];right=[x for x in s.shapes if x.left/Inches(1)>7 and 1<x.top/Inches(1)<6.8];top=max((x.top+x.height)/Inches(1) for x in right)+.08
 bottom=6.46;h=min(2.15,bottom-top);w=min(5.35,h*vs['width']/vs['height']);h=w*vs['height']/vs['width'];x=7.363+(5.35-w)/2
 assert h>1.5,(r['slide'],h)
 sh=s.shapes.add_picture(str(poster),Inches(x),Inches(top),Inches(w),Inches(h));sh.name=f"Added static video fallback {r['slide']}"
 sh=s.shapes.add_movie(str(embed),Inches(x),Inches(top),Inches(w),Inches(h),poster_frame_image=str(poster),mime_type='video/mp4');sh.name=f"Added academic video {r['slide']}"
 sh=s.shapes.add_textbox(Inches(7.363),Inches(6.5),Inches(5.35),Inches(.34));sh.name=f"Added video caption {r['slide']}";tf=sh.text_frame;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 rr=tf.paragraphs[0].add_run();rr.text=r['caption'];rr.font.name='Arial';rr.font.size=Pt(9);from pptx.dml.color import RGBColor;rr.font.color.rgb=RGBColor(0,0,0)
 tf=s.notes_slide.notes_text_frame
 for line in ['Additional original recording: '+r['caption'],'Article DOI: https://doi.org/'+r['doi'],'Original media URL: '+r['url'],'Credit: '+r['credit'],'Processing: '+r['processing'],'Teaching note: '+r['description']]+([full_reference] if r['slide']==4 else []):
  para=tf.add_paragraph();run=para.add_run();run.text=line;run.font.name='Arial'
 r.update(embed_path=str(embed),sha256=hashlib.sha256(embed.read_bytes()).hexdigest(),bytes=embed.stat().st_size,width=vs['width'],height=vs['height'],geometry={'x':x,'y':top,'w':w,'h':h})
p.save(out)
# Every old visible shape must remain byte-equivalent after canonical XML normalization.
ns={'p':'http://schemas.openxmlformats.org/presentationml/2006/main'}
with zipfile.ZipFile(src) as a,zipfile.ZipFile(out) as b:
 unchanged=[]
 for name in a.namelist():
  if name.startswith('ppt/media/'):
   assert a.read(name)==b.read(name),name;unchanged.append(name)
 for n in range(1,47):
  name=f'ppt/slides/slide{n}.xml';old=etree.fromstring(a.read(name)).find('p:cSld/p:spTree',ns);new=etree.fromstring(b.read(name)).find('p:cSld/p:spTree',ns)
  for o,v in zip(old,new):assert etree.tostring(o,method='c14n')==etree.tostring(v,method='c14n'),f'slide{n} existing shape changed'
  assert len(new)>=len(old)
 assert b.testzip() is None
 result=dict(slides=46,existing_shapes_unchanged=True,existing_media_unchanged=len(unchanged),added_videos=len(records),total_videos=18,bytes=out.stat().st_size,final_delivery_paused_for_three_aesthetic_uploads=True)
(B/'additional_videos/preparation_validation.json').write_text(json.dumps(result,indent=2)+'\n');(B/'additional_videos/ACADEMIC.json').write_text(json.dumps(records,indent=2)+'\n');print(json.dumps(result,indent=2))
