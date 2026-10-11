from pathlib import Path
import re,json,sys,zipfile,posixpath,hashlib
from lxml import etree as E
from pptx import Presentation
import pymupdf
B=Path(__file__).resolve().parent;R=B.parents[1];sys.path.insert(0,str(R/'tools'))
from style_rules import framing_problems,stats_problems,LIMITATION
PPT=B/'deliverables/Neuroethology_Lecture99_FA2026.pptx';PDF=Path('/workspace/l99-research/render/Neuroethology_Lecture99_FA2026.pdf')
p=Presentation(PPT);d=pymupdf.open(PDF);assert len(p.slides)==len(d)==46
errors=[];rows=[]
for i in range(1,45):
 sl=p.slides[i];sh=next(s for s in sl.shapes if s.name=='Body');ps=[q.text for q in sh.text_frame.paragraphs if q.text.strip()];assert len(ps)==3
 for j,s in enumerate(ps,1):
  count=len(re.findall(r'[.!?](?:\s|$)',s))
  if not 5<=count<=7:errors.append([i+1,j,'sentences',count])
  if framing_problems(s)or stats_problems(s)or LIMITATION.search(s):errors.append([i+1,j,'style',s])
 x,y,w,h=[v/914400*72 for v in [sh.left,sh.top,sh.width,sh.height]];lines=[]
 for b in d[i].get_text('dict')['blocks']:
  if b.get('type')!=0:continue
  for line in b['lines']:
   if abs(line['bbox'][0]-x)<2 and y-1<line['bbox'][1]<y+h+25 and max(t['size']for t in line['spans'])>=12.9:lines.append(line)
 normalize=lambda s:re.sub(r'\s+','',s).replace('–','-').replace('’',"'")
 counts=[];bottoms=[];index=0
 for j,text in enumerate(ps,1):
  target=normalize(text);actual='';count=0;lastbottom=None
  while index<len(lines)and len(actual)<len(target):
   line=lines[index];actual+=normalize(''.join(t['text']for t in line['spans']));count+=1;lastbottom=line['bbox'][3];index+=1
  counts.append(count)
  if lastbottom is not None:bottoms.append(lastbottom)
  if target!=actual:errors.append([i+1,j,'rendered text differs'])
 if index!=len(lines):errors.append([i+1,'unassigned body lines',len(lines)-index])
 if min(counts)<5:errors.append([i+1,'rendered lines',counts])
 if any(v>y+h+.5 for v in bottoms):errors.append([i+1,'text overflow',bottoms,y+h])
 rows.append({'slide':i+1,'sentences':[len(re.findall(r'[.!?](?:\s|$)',s))for s in ps],'rendered_lines':counts,'bottom_pt':max(bottoms)if bottoms else None,'body_bottom_pt':y+h})
report={'pages':46,'content_slides_checked':44,'paragraphs_checked':132,'minimum_rendered_lines':min(n for r in rows for n in r['rendered_lines']),'routine_limitations':0,'errors':errors,'slides':rows};(B/'rendered_validation.json').write_text(json.dumps(report,indent=2)+'\n')
packerrors=[];targets=[]
with zipfile.ZipFile(PPT) as z:
 assert z.testzip()is None;names=set(z.namelist())
 for name in names:
  if name.endswith('.xml')or name.endswith('.rels'):
   root=E.fromstring(z.read(name))
   if name.endswith('.rels'):
    base=posixpath.dirname(posixpath.dirname(name))
    for r in root:
     if r.get('TargetMode')=='External':continue
     t=r.get('Target');path=posixpath.normpath(posixpath.join(base,t))if not t.startswith('/')else t[1:]
     if path not in names:packerrors.append([name,t,'missing target'])
   if re.match(r'ppt/slides/slide\d+\.xml$',name):
    ns={'p':'http://schemas.openxmlformats.org/presentationml/2006/main'}
    ids=root.xpath('//p:cNvPr/@id',namespaces=ns)
    if len(ids)!=len(set(ids)):packerrors.append([name,'duplicate shape ids'])
  if name.startswith('ppt/media/')and name.endswith('.mp4'):targets.append((name,hashlib.sha256(z.read(name)).hexdigest(),len(z.read(name))))
 videos=json.loads((B/'VIDEOS.json').read_text());assert len(videos)==11
 for v in videos:
  found=[x for x in targets if x[1]==v['sha256']]
  if len(found)!=1:packerrors.append([v['key'],'embedded hash match count',len(found)])
summary_count=sum(1 for sl in p.slides for sh in sl.shapes if sh.has_text_frame and sh.text=='Key takeaways');assert summary_count==1
pack={'slides':46,'embedded_movies':len(targets),'distinct_recordings':len(set(x[1] for x in targets)),'movie_slides':[v['slide']for v in videos],'single_key_takeaways':summary_count,'zip_crc_errors':0,'xml_relationship_errors':packerrors,'bytes':PPT.stat().st_size,'sha256':hashlib.sha256(PPT.read_bytes()).hexdigest()};(B/'package_validation.json').write_text(json.dumps(pack,indent=2)+'\n')
print(json.dumps({k:v for k,v in report.items()if k!='slides'},indent=2));print(json.dumps(pack,indent=2))
if errors or packerrors:sys.exit(1)
