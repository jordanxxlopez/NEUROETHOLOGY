"""Apply measured text layout and embed original, credited article recordings."""
from pathlib import Path
import sys,json,importlib.util,subprocess,hashlib
from pptx import Presentation
from pptx.util import Inches,Pt
B=Path(__file__).resolve().parent;R=B.parents[1];sys.path.insert(0,str(R/'tools'))
from build_lecture import write_paras,picture_contain,textbox,plain
from theme_colors import darker
modpath=R/'lectures/L94/revision84/build_revision.py'
sp=importlib.util.spec_from_file_location('prior_layout',modpath);m=importlib.util.module_from_spec(sp);sp.loader.exec_module(m)
D=Path('/workspace/l99-research');M=D/'media';out=B/'deliverables';out.mkdir(exist_ok=True)
spec=json.loads((B/'lecture.json').read_text());prs=Presentation(D/'L99_Base.pptx');records=[];writing=[]
meta={f'j{i}':{'path':M/f'Jokura_video{i}.mp4','doi':'10.7554/eLife.108420','label':f'Jokura (2026), Video {i}.','url':f'https://cdn.elifesciences.org/articles/108420/elife-108420-video{i}.mp4','license':'CC BY','description':json.loads((B/'video_captions.json').read_text())[i-1]['caption']} for i in range(1,9)}
meta.update({'b1':{'path':M/'compatible/Burkhardt_movieS1.mp4','doi':'10.1126/science.ade5645','label':'Burkhardt (2023), Movie S1.','url':'https://www.science.org/doi/10.1126/science.ade5645#supplementary-materials','license':'Original article supplementary recording; authors credited','description':'3D reconstruction of the SNN, comb rows, sensory cells, mesogleal neurons and a tentacle from SBFSEM data of a 1-day-old cydippid.','processing':'Container remux from uploaded AVI to MP4, H.264 stream copied unchanged; no resizing or audio removal (source has no audio).'},
'p1':{'path':M/'compatible/Moroz_video1_clip.mp4','doi':'10.1038/nature13400','label':'Moroz (2014), Supplementary Video 1.','url':'https://www.nature.com/articles/nature13400#MOESM24','license':'Original supplementary recording, credited to Frazier Nivens / Ocean Imaging and Leonid L. Moroz / University of Florida','description':'Pleurobrachia bachei swimming, body motion and escape behavior.','processing':'Source interval 8–37 seconds; lossless H.264 conversion and lossless ALAC audio at original dimensions and rate. Intro and credits omitted from clip but credited in notes.'},
'p2':{'path':M/'compatible/Moroz_video2_clip.mp4','doi':'10.1038/nature13400','label':'Moroz (2014), Supplementary Video 2.','url':'https://www.nature.com/articles/nature13400#MOESM25','license':'Original supplementary recording, credited to Frazier Nivens / Ocean Imaging and Leonid L. Moroz / University of Florida','description':'Comparative ctenophore body forms, swimming and Beroe feeding on Bolinopsis.','processing':'Source interval 8–37 seconds; lossless H.264 conversion and lossless ALAC audio at original dimensions and rate. Intro and credits omitted from clip but credited in notes.'}})
short={'j1':'Cell organization within the statocyst.','j2':'Balancer cilia traced from their bases.','j3':'Syncytial soma and nuclei.','j4':'Sagittal ciliary arrest.','j5':'Sagittal renewed beating.','j6':'Tentacular ciliary arrest.','j7':'Tentacular renewed beating.','j8':'Tilt-induced asymmetric ciliary activity.','b1':'Surface nerve net and surrounding cells.','p1':'Sea-gooseberry swimming and escape.','p2':'Comparative swimming and feeding.'}
for n,sd in enumerate(spec['slides'],2):
 sl=prs.slides[n-1];body=next(sh for sh in sl.shapes if sh.name=='Body');pt,counts,need=m.fit(body,sd['body']);tf=body.text_frame;tf.clear();write_paras(tf,sd['body'],pt,'000000',bold_color=darker('435C70'),space_after=pt*.30)
 for p in tf.paragraphs:p.line_spacing=Pt(pt*1.15)
 writing.append({'slide':n,'font_pt':pt,'estimated_lines':counts,'sentences':[m.sentcount(p) for p in sd['body']],'height_in':need})
 if sd.get('video_key'):
  key=sd['video_key'];v=meta[key];fx=13.333-.62-5.35
  # Keep the published panel above the recording. Preserve every source pixel.
  for sh in list(sl.shapes):
   if sh.shape_type==13 or sh.name.startswith('Caption') or (sh.has_text_frame and sh.text==sd['figure']['caption']):
    el=sh._element;el.getparent().remove(el)
  picture_contain(sl,B/sd['figure']['path'],fx,1.4,5.35,2.24)
  _,ct=textbox(sl,fx,3.70,5.35,.53,'Figure caption');write_paras(ct,[sd['figure']['caption']],10,'000000',space_after=0)
  probe=json.loads(subprocess.check_output(['ffprobe','-v','error','-show_entries','stream=codec_type,codec_name,width,height,r_frame_rate:format=duration','-of','json',str(v['path'])]));vs=next(x for x in probe['streams'] if x['codec_type']=='video');ratio=vs['width']/vs['height'];h=2.10;w=min(5.35,h*ratio);h=w/ratio;x=fx+(5.35-w)/2
  poster=D/'posters'/f'{key}.png';fallback=sl.shapes.add_picture(str(poster),Inches(x),Inches(4.28),Inches(w),Inches(h));fallback.name=f'Static video fallback {key}';movie=sl.shapes.add_movie(str(v['path']),Inches(x),Inches(4.28),Inches(w),Inches(h),poster_frame_image=str(poster),mime_type='video/mp4');movie.name=f'Embedded movie {key}'
  _,vt=textbox(sl,fx,6.45,5.35,.32,'Video caption');write_paras(vt,[v['label']+' '+short[key]],9,'000000',space_after=0)
  text=sl.notes_slide.notes_text_frame.text;text+='\n\nVideo source: '+v['label']+' '+v['description']+'\nArticle DOI: https://doi.org/'+v['doi']+'\nOriginal media URL: '+v['url']+'\nCredit/license: '+v['license']+'\nProcessing: '+v.get('processing','Original MP4 embedded byte-for-byte; source-authored playback speed retained.')+'\nAudio present: '+str(any(x['codec_type']=='audio' for x in probe['streams']))
  sl.notes_slide.notes_text_frame.text=text
  records.append({**{k:str(z) if isinstance(z,Path) else z for k,z in v.items()},'slide':n,'key':key,'bytes':v['path'].stat().st_size,'sha256':hashlib.sha256(v['path'].read_bytes()).hexdigest(),'probe':probe,'caption':v['label']+' '+short[key]})
# Add source anatomy to molecular and muscle physiology slides.
for n,key in [(37,'s_net'),(39,'p_glutamate')]:
 sd=spec['slides'][n-2];sl=prs.slides[n-1];fx=13.333-.62-5.35
 for sh in list(sl.shapes):
  if sh.shape_type==13 or (sh.has_text_frame and sh.text==sd['figure']['caption']):el=sh._element;el.getparent().remove(el)
 assets=json.loads((B/'figure_sources.json').read_text());top=sd['figure'];bottom=assets[key]
 for f,y in [(top,1.4),(bottom,4.16)]:
  ph=picture_contain(sl,B/f['path'],fx,y,5.35,2.05);_,ct=textbox(sl,fx,y+2.08,5.35,.52,'Figure caption');write_paras(ct,[f['caption']],10,'000000',space_after=0)
 sl.notes_slide.notes_text_frame.text+='\nFigure: '+bottom['caption']+'\nFigure source: '+bottom['source_url']
for sl in prs.slides:
 for pp in sl.notes_slide.notes_text_frame.paragraphs:
  for rr in pp.runs:rr.font.name='Arial'
p=out/'Neuroethology_Lecture99_FA2026.pptx';prs.save(p)
(B/'VIDEOS.json').write_text(json.dumps(records,indent=2,ensure_ascii=False)+'\n');(B/'writing_validation.json').write_text(json.dumps(writing,indent=2)+'\n')
print('Saved',p,p.stat().st_size,'bytes;',len(prs.slides),'slides;',len(records),'embedded movies')
