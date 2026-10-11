"""Add ten verified clips to the released deck without rebuilding its content."""
from pathlib import Path
import json,hashlib,zipfile
from collections import defaultdict
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
B=Path(__file__).resolve().parent.parent;D=B/'additional_videos'
base=B/'Neuroethology_Lecture100_FA2026.pptx';out=B/'Neuroethology_Lecture100_FA2026_20_Videos.pptx'
new=json.loads((D/'media_sources.json').read_text())['added'];old=json.loads((B/'media_sources.json').read_text())['embedded'];p=Presentation(base)
assert len(p.slides)==46
original_text=[[sh.text for sh in sl.shapes if sh.has_text_frame]for sl in p.slides]
groups=defaultdict(list)
for m in new:groups[m['slide']].append(m)
for num,records in groups.items():
 slide=p.slides[num-1];existing_movies=[sh for sh in slide.shapes if sh.shape_type==16]
 original_figures=[sh for sh in slide.shapes if sh.shape_type==13 and not sh.name.startswith('Unedited article-recording frame:')]
 figure_captions=[sh for sh in slide.shapes if sh.name=='Figure caption']
 assert len(original_figures)==len(figure_captions)
 # Reflow the figure column only; preserve source content, labels and all text.
 for i,(pic,cap)in enumerate(zip(original_figures,figure_captions)):
  cell=5.7/len(original_figures);ratio=pic.width/pic.height;w=min(cell-.08,2.15*ratio);h=w/ratio
  pic.left=Inches(7.01+i*cell+(cell-w)/2);pic.top=Inches(1.35);pic.width=Inches(w);pic.height=Inches(h)
  cap.left=Inches(7.01+i*cell);cap.width=Inches(cell-.05);cap.top=Inches(1.35+h+.07);cap.height=Inches(.63)
  for para in cap.text_frame.paragraphs:
   for r in para.runs:r.font.size=Pt(9)
 side_by_side=num in (11,23,38,39) and not existing_movies and len(records)==1 and len(original_figures)==1
 if side_by_side:
  pic=original_figures[0];ratio=pic.width/pic.height;fw=min(2.4,4.45*ratio);fh=fw/ratio;pic.left=Inches(7.01);pic.top=Inches(1.35);pic.width=Inches(fw);pic.height=Inches(fh)
  cap=figure_captions[0];cap.left=Inches(7.01);cap.width=Inches(5.7);cap.top=Inches(6.02);cap.height=Inches(.46)
 total=len(existing_movies)+len(records);cell=5.7/total
 # Existing movie bytes, fallback frames, names and timing are retained.
 for i,movie in enumerate(existing_movies):
  ratio=movie.width/movie.height;w=min(cell-.06,2.35*ratio);h=w/ratio;left=7.01+i*cell+(cell-w)/2;top=4.23+(2.35-h)/2
  for sh in [movie]+[x for x in slide.shapes if x.name=='Unedited article-recording frame: '+movie.name.removesuffix(' (click to play)')]:
   sh.left=Inches(left);sh.top=Inches(top);sh.width=Inches(w);sh.height=Inches(h)
 for i,cap in enumerate([sh for sh in slide.shapes if sh.name=='Article recording credit']):
  cap.left=Inches(7.01+i*cell);cap.width=Inches(cell-.06)
 for offset,m in enumerate(records,len(existing_movies)):
  path=B/m['path'];poster=B/m['poster'];assert hashlib.sha256(path.read_bytes()).hexdigest()==m['sha256']
  w=min(cell-.06,2.35*m['display_aspect']);h=w/m['display_aspect'];left=7.01+offset*cell+(cell-w)/2;top=4.23+(2.35-h)/2
  if side_by_side:
   cw=5.7-fw-.2;w=min(cw,2.35*m['display_aspect']);h=w/m['display_aspect'];left=7.01+fw+.2+(cw-w)/2;top=3.4+(2.35-h)/2
  fallback=slide.shapes.add_picture(str(poster),Inches(left),Inches(top),Inches(w),Inches(h));fallback.name='Unedited additional-recording frame: '+m['caption']
  movie=slide.shapes.add_movie(str(path),Inches(left),Inches(top),Inches(w),Inches(h),poster_frame_image=str(poster),mime_type='video/mp4');movie.name=m['caption']+' (click to play)'
  cap=slide.shapes.add_textbox(Inches(7.01+offset*cell),Inches(6.58),Inches(cell-.06),Inches(.46));cap.name='Additional recording credit';tf=cap.text_frame;tf.word_wrap=True;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
  r=tf.paragraphs[0].add_run();r.text=m['caption'];r.font.name='Arial';r.font.size=Pt(8.5);r.font.color.rgb=RGBColor(0,0,0)
  notes=slide.notes_slide.notes_text_frame
  for line in ['',m['category']+': '+m['caption'],'• '+m['teaching_note'],'Original media URL: '+m['original_media_url'],'Source page: '+m['source_page'],'Credit: '+m['credit'],'License: '+m['license'],'Compatibility: '+m['conversion'],'Playback: click in desktop PowerPoint Slide Show. Media and relationships verified; interactive playback untested in this environment.']:
   notes.add_paragraph().text=line
  if m.get('doi'):notes.add_paragraph().text='Article DOI: https://doi.org/'+m['doi']
# Correct a verified bibliography typo only; lecture teaching text is untouched.
for slide in p.slides:
 for para in slide.notes_slide.notes_text_frame.paragraphs:
  if '10.1098/rspb.2012.1357' in para.text:
   for r in para.runs:r.text=r.text.replace('279:4550–4559','279:4559–4567')
# Every pre-existing visible text string survives in the same order.
for slide,before in zip(p.slides,original_text):
 after=[sh.text for sh in slide.shapes if sh.has_text_frame];assert after[:len(before)]==before
p.save(out)
with zipfile.ZipFile(out)as z:
 hashes={hashlib.sha256(z.read(n)).hexdigest()for n in z.namelist()if n.endswith('.mp4')}
 assert hashes=={m['sha256']for m in old+new}
 assert len(hashes)==20
 for m in old+new:
  rel=z.read(f"ppt/slides/_rels/slide{m['slide']}.xml.rels");assert b'/video' in rel and b'/media' in rel
  assert b'<p:timing>'in z.read(f"ppt/slides/slide{m['slide']}.xml")
assert len(Presentation(out).slides)==46
(D/'openxml_validation.json').write_text(json.dumps({'slides':46,'original_visible_text_preserved':True,'original_ten_recording_hashes_preserved':True,'additional_academic_recordings':7,'additional_natural_history_recordings':3,'total_embedded_recordings':20,'media_relationships':'passed','interactive_playback':'untested','pptx_bytes':out.stat().st_size},indent=2)+'\n')
print('Verified: 46 slides, original teaching text unchanged, all 20 recordings embedded.',flush=True)
