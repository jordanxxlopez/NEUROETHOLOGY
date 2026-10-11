from pathlib import Path
from pptx import Presentation
B=Path(__file__).resolve().parent.parent
p=Presentation(B/'Neuroethology_Lecture100_FA2026_20_Videos.pptx')
for slide in p.slides:
 for sh in list(slide.shapes):
  if sh.shape_type==16:sh._element.getparent().remove(sh._element)
 for rid,rel in list(slide.part.rels.items()):
  if rel.reltype.rsplit('/',1)[-1]in ('video','media','audio'):slide.part.drop_rel(rid)
 for el in list(slide._element):
  if el.tag.endswith('}timing'):slide._element.remove(el)
p.save(B/'Neuroethology_Lecture100_FA2026_20_Videos_static.pptx')
