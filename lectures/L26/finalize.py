import json,re
from pathlib import Path
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
H=Path(__file__).resolve().parent;root=H.parents[1];p=H/'Neuroethology_Lecture26_FA2026.pptx';prs=Presentation(p)
meta=next(x for x in json.loads((root/'course/schedule.json').read_text())['lectures']if x['n']==26)
sh=next(x for x in prs.slides[0].shapes if x.name=='Lecture title');sh.text_frame.clear();pa=sh.text_frame.paragraphs[0];pa.space_after=Pt(0);pa.line_spacing=1.08;r=pa.add_run();r.text=meta['title'];r.font.name='Arial';r.font.size=Pt(25);r.font.bold=True;r.font.color.rgb=RGBColor.from_string('283E4A');assert sh.text==meta['title']
spec=json.loads((H/'lecture.json').read_text());refs=json.loads((H/'references.json').read_text());nt=prs.slides[0].notes_slide.notes_text_frame;nt.text+='\n\n• The star-nosed mole uses 22 nasal rays to sample surfaces. Initial contacts recruit a broad sensory sheet; prey-directed movements bring the eleventh ray pair onto the target. The photograph is Gerhold et al. (2013), Figure 1A.\n\nReferences:\n'+refs['touch13']+'\n'+refs['fovea97']
prs.slides[-1].notes_slide.notes_text_frame.text='\n'.join('• '+i['lead']+'. '+i['text']for i in spec['takeaways']['items'])+'\n\nReferences:\n'+'\n'.join(spec['takeaways']['refs'])
for sl in prs.slides:
 for pa in sl.notes_slide.notes_text_frame.paragraphs:
  txt=pa.text;pa.clear()
  for chunk in re.split(r'(https://doi\.org/[^\s]+)',txt):
   if not chunk:continue
   run=pa.add_run();run.text=chunk;run.font.name='Arial';run.font.size=Pt(12)
   if chunk.startswith('https://doi.org/'):run.hyperlink.address=chunk
prs.save(p);print('Exact title; Arial notes; linked DOI references.')
