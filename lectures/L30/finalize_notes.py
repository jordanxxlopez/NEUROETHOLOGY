"""Apply Arial and plain, readable teaching text to the repository-built notes."""
from pathlib import Path
import sys
from pptx import Presentation
from pptx.util import Pt
ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT/'tools'))
from build_lecture import plain
path=Path(__file__).with_name('Neuroethology_Lecture30_FA2026.pptx')
prs=Presentation(path)
for slide in prs.slides:
 for paragraph in slide.notes_slide.notes_text_frame.paragraphs:
  paragraph.text=plain(paragraph.text)
  for run in paragraph.runs:
   run.font.name='Arial'
   run.font.size=Pt(12)
prs.save(path)
print('Applied Arial to all speaker notes and removed literal Markdown emphasis markers.')
