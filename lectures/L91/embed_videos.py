"""Add verified article recordings to the existing deck without rewriting slides."""
from pathlib import Path
import hashlib,json,subprocess
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from lxml import etree
ROOT=Path(__file__).resolve().parent
PPTX=ROOT/'Neuroethology_Lecture91_FA2026.pptx'
base='https://pmc-oa-opendata.s3.amazonaws.com/'
entries=[
 (20,'Rouse2016_Video1','Rouse et al. (2016), Supplementary Video 1. Ronan adjusts to phase and tempo changes.','10.3389/fnins.2016.00257',base+'PMC4891632.1/Video1.MP4','CC BY 4.0','Rouse, Cook, Large, and Reichmuth'),
 (41,'Ito2022_MovieS1','Ito et al. (2022), Movie S1. Rat head movements and simultaneous jerk traces.','10.1126/sciadv.abo7019',base+'PMC9651867.1/sciadv.abo7019_movies_s1_to_s3.zip','CC BY-NC 4.0','Ito, Shiramatsu, Ishida, Oshima, Magami, and Takahashi'),
 (42,'Ito2022_MovieS2','Ito et al. (2022), Movie S2. Beat-related movements in a standing rat.','10.1126/sciadv.abo7019',base+'PMC9651867.1/sciadv.abo7019_movies_s1_to_s3.zip','CC BY-NC 4.0','Ito, Shiramatsu, Ishida, Oshima, Magami, and Takahashi'),
 (43,'Ito2022_MovieS3','Ito et al. (2022), Movie S3. Human head movements and marker tracking.','10.1126/sciadv.abo7019',base+'PMC9651867.1/sciadv.abo7019_movies_s1_to_s3.zip','CC BY-NC 4.0','Ito, Shiramatsu, Ishida, Oshima, Magami, and Takahashi')]
additional=json.loads((ROOT/'additional_videos.json').read_text())
extra_by_stem={e['stem']:e for e in additional}
entries += [(e['slide'],e['stem'],e['caption'],e['doi'],base+e['original_key'],e['license'],e['credit']) for e in additional]
prs=Presentation(PPTX)
if any(s.shape_type==MSO_SHAPE_TYPE.MEDIA for sl in prs.slides for s in sl.shapes):
 raise SystemExit('Use the pre-video deck to avoid duplicate recordings.')
before=[[s.text for s in sl.shapes if s.has_text_frame] for sl in prs.slides]
original_images=[[hashlib.sha256(s.image.blob).hexdigest() for s in sl.shapes if s.shape_type==MSO_SHAPE_TYPE.PICTURE] for sl in prs.slides]
def fit(shape,x,y,w,h):
 ratio=shape.width/shape.height
 sw=min(w,h*ratio);sh=sw/ratio
 shape.left=Inches(x+(w-sw)/2);shape.top=Inches(y+(h-sh)/2);shape.width=Inches(sw);shape.height=Inches(sh)
def caption(sl,text,x,y,w,h):
 box=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
 tf=box.text_frame;tf.word_wrap=True;tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
 tf.text=text
 for p in tf.paragraphs:
  for r in p.runs:r.font.name='Arial';r.font.size=Pt(9);r.font.color.rgb=RGBColor(0,0,0)
 return box
sources=[]
for num,stem,cap,doi,url,license,credit in entries:
 sl=prs.slides[num-1];pic=next(s for s in sl.shapes if s.shape_type==MSO_SHAPE_TYPE.PICTURE)
 figcap=next(s for s in sl.shapes if s.has_text_frame and 'Fig.' in s.text)
 path=ROOT/'media'/f'{stem}.mp4';poster=ROOT/'posters'/f'{stem}.png'
 probe=json.loads(subprocess.check_output(['ffprobe','-v','error','-show_streams','-show_format','-of','json',str(path)]))
 v=next(s for s in probe['streams'] if s['codec_type']=='video');ratio=(float(v['display_aspect_ratio'].split(':')[0])/float(v['display_aspect_ratio'].split(':')[1])) if v.get('display_aspect_ratio') not in (None,'0:1') else v['width']/v['height']
 if num==43:
  fit(pic,7.11,1.4,2.6,2.5);figcap.left=Inches(7.11);figcap.top=Inches(4.0);figcap.width=Inches(2.6);figcap.height=Inches(.85)
  w=2.7;h=w/ratio;x=9.95;y=1.4
  caption(sl,cap,9.95,y+h+.12,2.7,.8)
 else:
  fit(pic,7.11,1.35,5.6,2.1);figcap.left=Inches(7.11);figcap.top=Inches(3.5);figcap.width=Inches(5.6);figcap.height=Inches(.42)
  h=2.35;w=h*ratio;x=7.11+(5.6-w)/2;y=4.02
  caption(sl,cap,7.11,6.48,5.6,.35)
 fallback=sl.shapes.add_picture(str(poster),Inches(x),Inches(y),width=Inches(w),height=Inches(h))
 fallback.name='Static video fallback'
 movie=sl.shapes.add_movie(str(path),Inches(x),Inches(y),Inches(w),Inches(h),poster_frame_image=str(poster),mime_type='video/mp4')
 for el in movie._element.xpath('.//p:cNvPr'):el.set('descr',cap+' Click to play the embedded original study recording.')
 if stem in extra_by_stem:
  e=extra_by_stem[stem]
  sl.notes_slide.notes_text_frame.text+='\nVideo comparison teaching: '+e['teaching']+'\nReferences:\n'+e['reference']+'\n'
 # PowerPoint supports full-screen video playback while preserving the native ratio.
 for el in sl._element.xpath('.//p:video'):el.set('fullScrn','1')
 sl.notes_slide.notes_text_frame.text+='\nEmbedded recording: '+cap+'\nSource URL: '+url+'\nArticle DOI: https://doi.org/'+doi+'\nCredit: '+credit+'. License: '+license+'.\nDiscuss the displayed behavior and compare it with the adjacent published measurements. Original timing and audio retained; poster is an unedited frame. Videos are H.264 MP4 copies for compatibility and delivery size, with native display proportions, original scene order, frame rate, duration, and audio retained. PCM audio in legacy MOV files is converted to AAC for compatibility. Desktop PowerPoint playback and Google Slides import were not tested.\n'
 sources.append(dict(slide=num,file=str(path.relative_to(ROOT)),caption=cap,doi=doi,original_url=url,credit=credit,license=license,width=v['width'],height=v['height'],duration_seconds=float(probe['format']['duration']),sha256=hashlib.sha256(path.read_bytes()).hexdigest(),display_aspect_ratio=v.get('display_aspect_ratio'),counts_toward_study_target=(stem!='Rouse2016_Video1')))
for i,sl in enumerate(prs.slides):
 assert [s.text for s in sl.shapes if s.has_text_frame][:len(before[i])]==before[i],f'Text changed on slide {i+1}'
 assert [hashlib.sha256(s.image.blob).hexdigest() for s in sl.shapes if s.shape_type==MSO_SHAPE_TYPE.PICTURE and s.name!='Static video fallback']==original_images[i]
prs.save(PPTX)
(ROOT/'media_sources.json').write_text(json.dumps(dict(embedded_video_count=len(sources),qualifying_non_ronan_video_count=sum(x['counts_toward_study_target'] for x in sources),required_count=10,ronan_study_recording_count=1,featured_2025_ronan_embedded=False,sources=sources),indent=2)+'\n')
print('Embedded eleven distinct recordings: ten non-Ronan recordings and one Ronan research recording. All existing visible text and original figure bytes preserved.')
