"""Embed 20 verified article recordings, retaining scientific figures and native proportions."""
from pathlib import Path
import json,hashlib,subprocess
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
R=Path(__file__).resolve().parent;p=R/'Neuroethology_Lecture104B_FA2026.pptx'
rows=json.loads((R/'media_preparation.json').read_text());byname={x['stem']:x for x in rows};prs=Presentation(p)
assert len(prs.slides)==46
# Static rendering found one 14-pt paragraph box overrun; keep the text at the repository 13-pt floor.
body=next(sh for sh in prs.slides[11].shapes if sh.name=='Body')
for pa in body.text_frame.paragraphs:
 pa.space_after=Pt(7.8)
 for ru in pa.runs:ru.font.size=Pt(13)
assert not any(x.shape_type==MSO_SHAPE_TYPE.MEDIA for s in prs.slides for x in s.shapes),'Rebuild before embedding.'
entries={}
for row in rows:
 caption=row['original_credit'].split(',')[0].split(' ')[0]+' et al. ('+str(json.loads((R/'reference_verification.json').read_text())[row['key']]['pubYear'])+'), '+row['original_identifier']+'. '+row['description']
 entries.setdefault(row['slide'],[]).append((row['stem'],caption))
assert sum(map(len,entries.values()))==20
spec=json.loads((R/'lecture.json').read_text());refs=json.loads((R/'reference_verification.json').read_text());sources=[]
def fit(pic,x,y,w,h):
 ratio=pic.width/pic.height;fw=min(w,h*ratio);fh=fw/ratio;pic.left=Inches(x+(w-fw)/2);pic.top=Inches(y+(h-fh)/2);pic.width=Inches(fw);pic.height=Inches(fh)
def cap(sl,text,x,y,w,h,name='Video caption'):
 box=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));box.name=name;tf=box.text_frame;tf.word_wrap=True;tf.margin_top=tf.margin_bottom=tf.margin_left=tf.margin_right=0;tf.text=text
 for para in tf.paragraphs:
  for run in para.runs:run.font.name='Arial';run.font.size=Pt(9);run.font.color.rgb=RGBColor(0,0,0)
 return box
for n,items in entries.items():
 sl=prs.slides[n-1];before=[sh.text for sh in sl.shapes if sh.has_text_frame];pics=[sh for sh in sl.shapes if sh.shape_type==MSO_SHAPE_TYPE.PICTURE];captions=[sh for sh in sl.shapes if sh.name=='Figure caption']
 # The source figure remains distinct from its companion image and actual movie.
 if len(pics)>1:
  fit(pics[0],7.11,1.34,3.92,2.40);fit(pics[1],11.18,1.34,1.52,2.40)
  for j,(x,w) in enumerate([(7.11,3.92),(11.18,1.52)]):
   captions[j].left=Inches(x);captions[j].top=Inches(3.77);captions[j].width=Inches(w);captions[j].height=Inches(.61)
   for pa in captions[j].text_frame.paragraphs:
    for ru in pa.runs:ru.font.size=Pt(8)
 else:
  fit(pics[0],7.11,1.34,5.59,2.50);captions[0].left=Inches(7.11);captions[0].top=Inches(3.87);captions[0].width=Inches(5.59);captions[0].height=Inches(.48)
 count=len(items);slotw=(5.59-.18*(count-1))/count
 for j,(stem,caption) in enumerate(items):
  row=dict(byname[stem]);file=R/row['delivery_file'];post=R/row['poster'];probe=json.loads(subprocess.check_output(['ffprobe','-v','error','-show_streams','-of','json',str(file)]));v=next(x for x in probe['streams'] if x['codec_type']=='video');dar=v.get('display_aspect_ratio');ratio=float(dar.split(':')[0])/float(dar.split(':')[1]) if dar and dar!='0:1' else v['width']/v['height']
  maxh=2.10;w=min(slotw,maxh*ratio);h=w/ratio;x=7.11+j*(slotw+.18)+(slotw-w)/2;y=4.48+(maxh-h)/2
  fallback=sl.shapes.add_picture(str(post),Inches(x),Inches(y),width=Inches(w),height=Inches(h));fallback.name='Static video fallback'
  movie=sl.shapes.add_movie(str(file),Inches(x),Inches(y),Inches(w),Inches(h),poster_frame_image=str(post),mime_type='video/mp4');movie.name=stem
  for el in movie._element.xpath('.//p:cNvPr'):el.set('descr',caption+' Click to play the embedded original recording.')
  for el in sl._element.xpath('.//p:video'):el.set('fullScrn','1')
  cap(sl,caption,7.11+j*(slotw+.18),6.57,slotw,.38)
  doi=row['article_doi'];ref=next(x for x in refs.values() if x['doi'].lower()==doi.lower());credit=ref.get('authorString','Original study authors')
  license='CC BY 4.0' if row['key'] in ('Rat2018','PVT2025','Care2026') else 'Original publisher recording; no open license stated in author manuscript'
  row.update({'slide':n,'caption':caption,'credit':credit,'license':license,'embedded':True,'native_aspect_ratio':ratio,'genre':row.get('genre','original experimental recording')});sources.append(row)
  sl.notes_slide.notes_text_frame.text+='\n\nEmbedded video: '+caption+'\nExact source identifier: '+row['original_identifier']+'\nArticle DOI: https://doi.org/'+doi+'\nOriginal media URL: '+row['download_url']+'\nCredit: '+credit+' License: '+license+'.\n'+row['conversion']+'. Poster is a source frame. Playback is full screen where the viewer supports it. Desktop PowerPoint GUI playback and Google Slides import remain untested. Google Slides may require separate Drive media insertion.\n'
 assert [sh.text for sh in sl.shapes if sh.has_text_frame][:len(before)]==before,f'Slide {n} teaching text changed'
prs.save(p);(R/'media_sources.json').write_text(json.dumps({'embedded_video_count':len(sources),'distinct_original_recordings':len({x['original_sha256'] for x in sources}),'experimental_recordings':15,'author_figure_presentations':0,'sources':sources},indent=2)+'\n');print('Embedded 20 clips on',len(entries),'slides. PPTX bytes:',p.stat().st_size)
