"""Build with repository tools, then position the unusually long exact title."""
from pathlib import Path
import subprocess,sys
from pptx import Presentation
from pptx.util import Inches
D=Path(__file__).resolve().parent;ROOT=D.parents[1];ppt=D/'Neuroethology_Lecture60_FA2026.pptx'
subprocess.run([sys.executable,str(ROOT/'tools/build_lecture.py'),str(D/'lecture.json')],check=True)
r=Presentation(ppt)
for x in r.slides[0].shapes:
 if not x.has_text_frame:continue
 if x.text.startswith('Seahorses:'):x.top=Inches(2.15)
 elif x.text=='Lecture 60':x.top=Inches(1.65)
 elif x.text=='Date TBD':x.top=Inches(5.95)
 elif 'Professor of Neurobiology' in x.text:x.top=Inches(6.65);x.height=Inches(.65)
r.save(ppt)
subprocess.run([sys.executable,str(ROOT/'tools/check_lecture.py'),str(ppt),'--lecture','60'],check=True)
