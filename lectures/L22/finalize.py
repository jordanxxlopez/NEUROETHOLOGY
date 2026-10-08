"""Preserve the character-exact schedule title and format complete notes."""
import json,re
from pathlib import Path
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
H=Path(__file__).resolve().parent;root=H.parents[1];p=H/'Neuroethology_Lecture22_FA2026.pptx';prs=Presentation(p)
meta=next(x for x in json.loads((root/'course/schedule.json').read_text())['lectures']if x['n']==22)
sh=next(x for x in prs.slides[0].shapes if x.name=='Lecture title');sh.text_frame.clear();pa=sh.text_frame.paragraphs[0];pa.space_after=Pt(0);pa.line_spacing=1.08;r=pa.add_run();r.text=meta['title'];r.font.name='Arial';r.font.size=Pt(24);r.font.bold=True;r.font.color.rgb=RGBColor.from_string('FFFFFF');assert sh.text==meta['title']
spec=json.loads((H/'lecture.json').read_text());refs=json.loads((H/'references.json').read_text())
nt=prs.slides[0].notes_slide.notes_text_frame;nt.text+='\n\n• Eigenmannia virescens produces a continuous wave-type electric organ discharge. The title image is the original published animal figure from Heiligenberg (1980).\n\nReferences:\n'+refs['control80']
prs.slides[-1].notes_slide.notes_text_frame.text='\n'.join('• '+i['lead']+'. '+i['text']for i in spec['takeaways']['items'])+'\n\nReferences:\n'+'\n'.join(spec['takeaways']['refs'])
for sl in prs.slides:
 tf=sl.notes_slide.notes_text_frame
 for pa in tf.paragraphs:
  txt=pa.text;pa.clear()
  for chunk in re.split(r'(https://doi\.org/[^\s]+)',txt):
   if not chunk:continue
   run=pa.add_run();run.text=chunk;run.font.name='Arial';run.font.size=Pt(12)
   if chunk.startswith('https://doi.org/'):run.hyperlink.address=chunk
prs.save(p);print('Exact title retained; all notes Arial; DOI links clickable.')
