"""Preserve the exact schedule title; format notes and link their source DOIs."""
import json,re
from pathlib import Path
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
H=Path(__file__).resolve().parent;ROOT=H.parents[1];p=H/'Neuroethology_Lecture29_FA2026.pptx';prs=Presentation(p)
spec=json.loads((H/'lecture.json').read_text());meta=next(x for x in json.loads((ROOT/'course/schedule.json').read_text())['lectures']if x['n']==29)
sh=next(x for x in prs.slides[0].shapes if x.name=='Lecture title');sh.text_frame.clear();pa=sh.text_frame.paragraphs[0];pa.space_after=Pt(0);pa.line_spacing=1.08;r=pa.add_run();r.text=meta['title'];r.font.name='Arial';r.font.size=Pt(25);r.font.bold=True;r.font.color.rgb=RGBColor.from_string('3B4149')
nt=prs.slides[0].notes_slide.notes_text_frame;nt.text='• '+meta['title']+'\n• Honeybee foragers communicate a resource vector on the comb. The study animal is photographed within Dong et al. (2023), Figure 1A; the complete published panel retains its labels and dance-path annotation.\n• Direction, visually calibrated distance, time compensation and terminal search are experimentally separable components of recruitment.\n\nFigure: '+spec['title_image']['caption']+'\nFigure source: '+spec['title_image']['source_url']+'\n\nReferences:\n'+'\n'.join(spec['title_refs'])
prs.slides[-1].notes_slide.notes_text_frame.text='\n'.join('• '+i['lead']+'. '+i['text']for i in spec['takeaways']['items'])+'\n\nReferences:\n'+'\n'.join(spec['takeaways']['refs'])
for sl in prs.slides:
 for pa in sl.notes_slide.notes_text_frame.paragraphs:
  txt=pa.text;pa.clear()
  for chunk in re.split(r'(https://doi\.org/[^\s]+)',txt):
   if not chunk:continue
   run=pa.add_run();run.text=chunk;run.font.name='Arial';run.font.size=Pt(12)
   if chunk.startswith('https://doi.org/'):run.hyperlink.address=chunk
prs.save(p)
