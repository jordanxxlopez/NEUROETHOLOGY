"""Preserve exact schedule title; use Arial and linked DOIs in teaching notes."""
import json,re
from pathlib import Path
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
H=Path(__file__).resolve().parent;ROOT=H.parents[1];p=H/'Neuroethology_Lecture36_FA2026.pptx';prs=Presentation(p)
spec=json.loads((H/'lecture.json').read_text());meta=next(x for x in json.loads((ROOT/'course/schedule.json').read_text())['lectures']if x['n']==36)
sh=next(x for x in prs.slides[0].shapes if x.name=='Lecture title');sh.text_frame.clear();pa=sh.text_frame.paragraphs[0];pa.space_after=Pt(0);pa.line_spacing=1.08;r=pa.add_run();r.text=meta['title'];r.font.name='Arial';r.font.size=Pt(25);r.font.bold=True;r.font.color.rgb=RGBColor.from_string('60353B')
nt=prs.slides[0].notes_slide.notes_text_frame;nt.text='• '+meta['title']+'\n• Aplysia californica provides identifiable sensory and motor neurons connecting tactile stimulation to defensive withdrawal. Habituation and sensitization permit behavioral, synaptic, and molecular interventions in related preparations.\n• The animal photograph is credited to Jeff Young, CC BY 4.0. https://www.inaturalist.org/photos/746343526\n\nReferences:\n'+'\n'.join(spec['title_refs'])
prs.slides[-1].notes_slide.notes_text_frame.text='\n'.join('• '+i['lead']+'. '+i['text']for i in spec['takeaways']['items'])+'\n\nReferences:\n'+'\n'.join(spec['takeaways']['refs'])
for sl in prs.slides:
 for pa in sl.notes_slide.notes_text_frame.paragraphs:
  txt=pa.text;pa.clear()
  for chunk in re.split(r'(https://doi\.org/[^\s]+)',txt):
   if not chunk:continue
   run=pa.add_run();run.text=chunk;run.font.name='Arial';run.font.size=Pt(12)
   if chunk.startswith('https://doi.org/'):run.hyperlink.address=chunk
prs.save(p)
