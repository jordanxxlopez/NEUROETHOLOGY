from pathlib import Path
import json,re
from pptx import Presentation
from pptx.util import Pt
BASE=Path(__file__).resolve().parent;p=BASE/'Neuroethology_Lecture41_FA2026.pptx';prs=Presentation(p);spec=json.loads((BASE/'lecture.json').read_text())
last=prs.slides[-1];last.notes_slide.notes_text_frame.text='\n'.join('• '+x['lead']+' '+x['text'] for x in spec['takeaways']['items'])+'\n\nReferences:\n'+'\n'.join(spec['takeaways']['refs'])
for par in last.notes_slide.notes_text_frame.paragraphs:
 for r in par.runs:
  r.font.name='Arial'
  if r.text.startswith('https://doi.org/'):r.hyperlink.address=r.text

# DOI text remains intact, with actual clickable hyperlinks in the notes.
for sl in prs.slides:
 tf=sl.notes_slide.notes_text_frame
 for par in tf.paragraphs:
  raw=par.text
  if 'https://doi.org/' not in raw:continue
  par.clear();pos=0
  for m in re.finditer(r'https://doi\.org/[^\s]+',raw):
   r=par.add_run();r.text=raw[pos:m.start()];r.font.name='Arial';r=par.add_run();r.text=m.group();r.font.name='Arial';r.hyperlink.address=m.group();pos=m.end()
  r=par.add_run();r.text=raw[pos:];r.font.name='Arial'
 for par in tf.paragraphs:
  for r in par.runs:r.font.name='Arial'
# Retain the exact schedule string as one title paragraph.
assert next(x for x in prs.slides[0].shapes if x.name=='Lecture title').text==json.loads((BASE.parents[1]/'course/schedule.json').read_text())['lectures'][-1]['title']
prs.save(p)
