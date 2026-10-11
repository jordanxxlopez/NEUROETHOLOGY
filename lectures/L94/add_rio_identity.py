"""Replace the slide-35 Joly movie and add Rio on slide 32, retaining source quality."""
from pathlib import Path
import json,hashlib,zipfile,copy
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
B=Path(__file__).resolve().parent;src=B/'deliverables/Neuroethology_Lecture94_FA2026_Paragraph3_Revised.pptx';p=Presentation(src)
f=B/'media/Rio_identity.mp4';poster=B/'media/Rio_identity_poster.png';source_bytes=f.read_bytes();poster_bytes=poster.read_bytes()
ref='Cook P, Reichmuth C, Hanke FD (2021). The Mind of a Sea Lion. In: Campagna C, Harcourt R (eds), Ethology and Behavioral Ecology of Otariids and the Odobenid. Springer, Cham, pp. 323–345. https://doi.org/10.1007/978-3-030-59184-7_17'
url='https://link.springer.com/chapter/10.1007/978-3-030-59184-7_17';credit='Cook, Reichmuth and Hanke; video supplied by the instructor. Publisher identifies Rio in Supplementary Video 17.2; the opening title card identifies Video Supplement 2.'
original_text=[[s.text for s in sl.shapes if s.has_text_frame]for sl in p.slides];old_movies=[(i,s.shape_id,s.element.xml)for i,sl in enumerate(p.slides)for s in sl.shapes if s.shape_type==16 and i!=34]
sl=p.slides[34];movie=next(s for s in sl.shapes if s.shape_type==16);old_credit=next(s for s in sl.shapes if s.name=='Article recording credit');fallback=next(s for s in sl.shapes if s.name.startswith('Unedited article-recording frame:'))
video_rel=next(r for r in sl.part.rels.values()if r.reltype.endswith('/video'));media_part=video_rel.target_part;old_video_part=str(media_part.partname).lstrip('/');old_video_sha=hashlib.sha256(media_part.blob).hexdigest();media_part._blob=source_bytes
image_rel=sl.part.rels[fallback._pic.blipFill.blip.rEmbed];image_part=image_rel.target_part;old_poster_part=str(image_part.partname).lstrip('/');image_part._blob=poster_bytes
w=2.35*16/9;x=7.01+(5.7-w)/2
for sh in [movie,fallback]:sh.left=Inches(x);sh.width=Inches(w)
caption35='Cook et al. (2021), Supplementary Video 17.2. Rio identity matching, comparison task.'
movie.name=caption35+' (click to play)';fallback.name='Unedited Rio identity-matching source frame';old_credit.text_frame.paragraphs[0].runs[0].text='Click to play. '+caption35
for r in old_credit.text_frame.paragraphs[0].runs[1:]:r.text=''
notes=sl.notes_slide.notes_text_frame;pars=list(notes.paragraphs);start=next(i for i,par in enumerate(pars)if par.text.startswith('Article-provided recording: Joly'))
for par in pars[start:]:par._p.getparent().remove(par._p)
sl32=p.slides[31];pic=next(s for s in sl32.shapes if s.shape_type==13);cap=next(s for s in sl32.shapes if s.name=='Figure caption');ratio=pic.width/pic.height;fw=min(5.7,2.5*ratio);fh=fw/ratio
pic.left=Inches(7.01+(5.7-fw)/2);pic.top=Inches(1.35);pic.width=Inches(fw);pic.height=Inches(fh);cap.left=Inches(7.01);cap.top=Inches(1.35+fh+.05);cap.width=Inches(5.7);cap.height=Inches(.43)
for par in cap.text_frame.paragraphs:
 for r in par.runs:r.font.size=Pt(9)
fall=sl32.shapes.add_picture(str(poster),Inches(x),Inches(4.23),Inches(w),Inches(2.35));fall.name='Unedited Rio identity-matching source frame'
caption32='Cook et al. (2021), Supplementary Video 17.2. Rio chooses the shape matching the central sample.'
mov=sl32.shapes.add_movie(str(f),Inches(x),Inches(4.23),Inches(w),Inches(2.35),poster_frame_image=str(poster),mime_type='video/mp4');mov.name=caption32+' (click to play)'
box=sl32.shapes.add_textbox(Inches(7.01),Inches(6.62),Inches(5.7),Inches(.4));box.name='Article recording credit';tf=box.text_frame;tf.word_wrap=True;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;r=tf.paragraphs[0].add_run();r.text='Click to play. '+caption32;r.font.name='Arial';r.font.size=Pt(9);r.font.color.rgb=RGBColor(0,0,0)
for number,caption in [(32,caption32),(35,caption35)]:
 notes=p.slides[number-1].notes_slide.notes_text_frame
 teaching='Rio inspects a sample shape in the center position. When the two side choices become available, she contacts the shape matching that sample. The recording contains two identity-matching choices, first toward the right and then toward the left. The publisher identifies a tone and fish as consequences of correct choices.'
 distinction='This is an identity-matching comparison, not footage of harbor-seal reversal learning or a reward reassignment. Identity matching supplies a sample-dependent correct choice; reversal changes an acquired stimulus–reward assignment. The Erdsack study and its apparatus figure remain the sources for the reversal findings on this slide.'if number==35 else'The footage records identity-matching performance. The published chapter cites Kastak and Schusterman (1994), the sea-lion identity-transfer study discussed on this slide. The publisher does not identify these filmed trials as novel-exemplar or first-trial transfer tests, and their filming date is not stated.'
 for t in ['', 'Article-provided recording: '+caption,'• '+teaching,distinction,'References: '+ref,'Publisher supplementary entry: '+url+'#MOESM2','Original source player: https://e.video-cdn.net/watch?video-id=6t9bp_2EXpiiKeBUT4Rt3w','Credit: '+credit,'License: © 2021 Springer Nature Switzerland AG; no separate open license stated for the recording. Instructor supplied this copy.','Source filename: Rio_identity.mp4. Original uploaded copy retained byte-for-byte: 1280 × 720, 25 frames/s, H.264, AAC stereo, 41.006 s. No cropping, trimming, re-encoding, compression or resolution reduction.','Playback: click the poster in desktop PowerPoint Slide Show. Embedded media, timing and source decoding were checked; interactive PowerPoint playback is untested.']:notes.add_paragraph().text=t
for i,sl in enumerate(p.slides):
 texts=[s.text for s in sl.shapes if s.has_text_frame]
 if i==34:
  old=list(original_text[i]);old[-1]=texts[-1];assert old==texts
 else:assert texts[:len(original_text[i])]==original_text[i]
for i,sid,xml in old_movies:assert next(s for s in p.slides[i].shapes if s.shape_id==sid).element.xml==xml
out=B/'deliverables/Neuroethology_Lecture94_FA2026_Rio_Revised.pptx';p.save(out)
with zipfile.ZipFile(src)as z,zipfile.ZipFile(out)as n:
 for name in z.namelist():
  if name.startswith('ppt/media/')and name not in [old_video_part,old_poster_part]:assert z.read(name)==n.read(name),name
 media={name:hashlib.sha256(n.read(name)).hexdigest()for name in n.namelist()if name.startswith('ppt/media/')and Path(name).suffix in ['.mp4','.mpg','.mov']}
 assert old_video_sha not in media.values()and hashlib.sha256(source_bytes).hexdigest()in media.values()
assert len(p.slides)==46 and sum(s.shape_type==16 for sl in p.slides for s in sl.shapes)==33
meta={'slides':[32,35],'source_filename':f.name,'path':str(f.relative_to(B)),'poster':str(poster.relative_to(B)),'doi':'10.1007/978-3-030-59184-7_17','reference':ref,'publisher_url':url+'#MOESM2','publisher_identifier':'Supplementary Video 17.2','on_video_identifier':'Video Supplement 2','publisher_description':'Sea lion Rio performing an identity matching-to-sample task. Correct responses (selecting the choice stimulus that matches the shape in the center position) are marked by a tone and reinforced with fish.','credit':credit,'license':'© 2021 Springer Nature Switzerland AG; separate video reuse license not stated. Instructor supplied copy.','video_sha256':hashlib.sha256(source_bytes).hexdigest(),'duration_s':41.006009,'width':1280,'height':720,'audio_present':True,'editing':'Uploaded video retained byte-for-byte; used on exactly two slides without temporal or spatial changes.','replaced_recording':{'slide':35,'source':'Joly et al. (2014), Movie S1','old_video_sha256':old_video_sha},'embedded_movie_shapes':33,'output_file':str(out.relative_to(B)),'bytes':out.stat().st_size,'sha256':hashlib.sha256(out.read_bytes()).hexdigest()}
(B/'rio_recording.json').write_text(json.dumps(meta,indent=2,ensure_ascii=False)+'\n');print(out,out.stat().st_size,'bytes; 33 video shapes, Rio used twice; lecture text and all other media unchanged.')
