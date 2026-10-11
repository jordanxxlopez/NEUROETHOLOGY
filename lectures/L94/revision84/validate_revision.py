"""Sentence, rendered-line, text-bound and preserved-shape checks."""
from pathlib import Path
import re,json,sys,zipfile
from lxml import etree as E
from pptx import Presentation
import pymupdf as fitz
sys.path.insert(0,str(Path(__file__).resolve().parents[3]/'tools'))
from style_rules import framing_problems,stats_problems,LIMITATION
B=Path(__file__).resolve().parent
BASE=B.parent/'deliverables/Neuroethology_Lecture94_FA2026_70slides_SingleSummary.pptx'
NEW=B.parent/'deliverables/Neuroethology_Lecture94_FA2026_84slides_Revised.pptx'
PDF=Path(sys.argv[1])
p=Presentation(NEW);old=Presentation(BASE);d=fitz.open(PDF)
assert len(p.slides)==len(d)==84
errors=[];rows=[]
def canon(e):return E.tostring(e,method='c14n')
for i in range(1,83):
    sl=p.slides[i];sh=next(s for s in sl.shapes if s.name=='Body')
    ps=[q.text for q in sh.text_frame.paragraphs if q.text.strip()]
    assert len(ps)==3
    for j,s in enumerate(ps,1):
        count=len(re.findall(r'[.!?](?:\s|$)',s))
        if not 5<=count<=7:errors.append([i+1,j,'sentences',count])
        if framing_problems(s)or stats_problems(s)or LIMITATION.search(s):errors.append([i+1,j,'style',s])
    x,y,w,h=[v/914400*72 for v in [sh.left,sh.top,sh.width,sh.height]]
    lines=[]
    for b in d[i].get_text('dict')['blocks']:
        if b.get('type')!=0:continue
        for line in b['lines']:
            if abs(line['bbox'][0]-x)<2 and y-1<line['bbox'][1]<y+h+25 and max(t['size']for t in line['spans'])>=12.9:
                lines.append(line)
    normalize=lambda s:re.sub(r'\s+','',s).replace('–','-').replace('’',"'")
    counts=[];bottoms=[];index=0
    for j,text in enumerate(ps,1):
        target=normalize(text);actual='';count=0;lastbottom=None
        while index<len(lines)and len(actual)<len(target):
            line=lines[index];actual+=normalize(''.join(t['text']for t in line['spans']));count+=1
            lastbottom=line['bbox'][3];index+=1
        counts.append(count)
        if lastbottom is not None:bottoms.append(lastbottom)
        if target!=actual:errors.append([i+1,j,'rendered text differs'])
    if index!=len(lines):errors.append([i+1,'unassigned body lines',len(lines)-index])
    if min(counts)<5:errors.append([i+1,'rendered lines',counts])
    if any(v>y+h+.5 for v in bottoms):errors.append([i+1,'text overflow',bottoms,y+h])
    rows.append({'slide':i+1,'sentences':[len(re.findall(r'[.!?](?:\s|$)',s))for s in ps],'rendered_lines':counts,'bottom_pt':max(bottoms)if bottoms else None,'body_bottom_pt':y+h})
# Compare every pre-existing shape. Only body and final summary footer may differ.
checked=0
for i,os in enumerate(old.slides):
    ns=p.slides[i if i<69 else 83]
    assert len(os.shapes)==len(ns.shapes),(i+1,len(os.shapes),len(ns.shapes))
    for a,b in zip(os.shapes,ns.shapes):
        if (i not in [0,69]and a.name=='Body')or(i==69 and a.name=='Slide number'):continue
        if canon(a._element)!=canon(b._element):errors.append([i+1,'preserved shape differs',a.name])
        checked+=1
report={'pages':len(d),'content_slides_checked':82,'paragraphs_checked':246,'minimum_rendered_lines':min(n for r in rows for n in r['rendered_lines']),'existing_shapes_preserved':checked,'routine_limitations':0,'errors':errors,'slides':rows}
(B/'rendered_validation.json').write_text(json.dumps(report,indent=2)+'\n')
print(json.dumps({k:v for k,v in report.items()if k!='slides'},indent=2))
if errors:sys.exit(1)
