"""Add seven native article recordings to the preserved 25-video Lecture 94.
Run after add_uploaded_excerpts.py. Original source codecs/bytes are preserved.
"""
from pathlib import Path
import json,hashlib,zipfile,urllib.request
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
B=Path(__file__).resolve().parent
m=json.loads((B/'media_sources.json').read_text())
src=B/'deliverables/Neuroethology_Lecture94_FA2026.pptx'
p=Presentation(src)
assert sum(s.shape_type==16 for sl in p.slides for s in sl.shapes)==25
texts=[[s.text for s in sl.shapes if s.has_text_frame]for sl in p.slides]
movies=[(i,s.shape_id,s.element.xml)for i,sl in enumerate(p.slides)for s in sl.shapes if s.shape_type==16]
for c in m['late_slide_recordings']:
 f=B/c['path'];poster=B/c['poster']
 if not f.exists():
  f.parent.mkdir(parents=True,exist_ok=True)
  with urllib.request.urlopen(c['original_media_url'])as r,f.open('wb')as o:
   while chunk:=r.read(1024*1024):o.write(chunk)
 assert hashlib.sha256(f.read_bytes()).hexdigest()==c['sha256']==c['original_sha256']
 sl=p.slides[c['slide']-1];assert not any(s.shape_type==16 for s in sl.shapes)
 pics=[s for s in sl.shapes if s.shape_type==13];caps=[s for s in sl.shapes if s.name.startswith('Figure caption')]
 assert len(pics)==len(caps)
 for j,(pic,cap)in enumerate(zip(pics,caps)):
  ratio=pic.width/pic.height;cell=5.7 if len(pics)==1 else 2.75;left=7.01+j*(2.95 if len(pics)>1 else 0)
  w=min(cell,2.5*ratio);h=w/ratio
  pic.left=Inches(left+(cell-w)/2);pic.top=Inches(1.35);pic.width=Inches(w);pic.height=Inches(h)
  cap.left=Inches(left);cap.width=Inches(cell);cap.top=Inches(1.35+h+.05);cap.height=Inches(.48)
  for para in cap.text_frame.paragraphs:
   for r in para.runs:r.font.size=Pt(8 if len(pics)>1 else 9)
 w=min(5.7,2.35*c['display_aspect']);h=w/c['display_aspect'];x=7.01+(5.7-w)/2;top=4.23
 fallback=sl.shapes.add_picture(str(poster),Inches(x),Inches(top),Inches(w),Inches(h));fallback.name='Unedited comparative article recording source frame'
 mime={'.mp4':'video/mp4','.mpg':'video/mpeg','.mov':'video/quicktime'}[f.suffix]
 movie=sl.shapes.add_movie(str(f),Inches(x),Inches(top),Inches(w),Inches(h),poster_frame_image=str(poster),mime_type=mime);movie.name=c['caption']+' (click to play)'
 box=sl.shapes.add_textbox(Inches(7.01),Inches(top+h+.06),Inches(5.7),Inches(.46));box.name='Comparative article video credit'
 tf=box.text_frame;tf.word_wrap=True;tf.margin_top=tf.margin_bottom=tf.margin_left=tf.margin_right=0
 r=tf.paragraphs[0].add_run();r.text='Click to play. Comparative recording: '+c['caption'];r.font.name='Arial';r.font.size=Pt(9);r.font.color.rgb=RGBColor(0,0,0)
 notes=sl.notes_slide.notes_text_frame
 for t in ['', 'Added comparative recording: '+c['caption'],'• '+c['teaching_note'],'References: '+c['reference'],'Original article recording: '+c['original_media_url'],'Credit: '+c['credit'],'License: '+c['license'],'Original filename: '+c['original_filename'],c['conversion'],c['playback_note']]:notes.add_paragraph().text=t
 if c.get('source_decoder_note'):notes.add_paragraph().text=c['source_decoder_note']
for i,sl in enumerate(p.slides):assert [s.text for s in sl.shapes if s.has_text_frame][:len(texts[i])]==texts[i]
for i,sid,xml in movies:assert next(s for s in p.slides[i].shapes if s.shape_id==sid).element.xml==xml
out=B/'deliverables/Neuroethology_Lecture94_FA2026_32videos.pptx';p.save(out)
with zipfile.ZipFile(src)as old,zipfile.ZipFile(out)as new:
 for n in old.namelist():
  if n.startswith('ppt/media/'):assert old.read(n)==new.read(n)
 expected={c['sha256']for key in ['embedded','eve_excerpts','additional_uploaded_excerpts','late_slide_recordings']for c in m[key]}
 actual={hashlib.sha256(new.read(n)).hexdigest()for n in new.namelist()if n.startswith('ppt/media/')and Path(n).suffix.lower()in {'.mp4','.mpg','.mov'}}
 assert expected==actual
 for c in m['late_slide_recordings']:
  assert b'<p:timing>'in new.read(f"ppt/slides/slide{c['slide']}.xml")
  assert b'/video'in new.read(f"ppt/slides/_rels/slide{c['slide']}.xml.rels")
assert len(p.slides)==46 and sum(s.shape_type==16 for sl in p.slides for s in sl.shapes)==32
print(out, out.stat().st_size,'bytes; 32 embedded movies; original text, 25 movies and source-image bytes preserved.')
