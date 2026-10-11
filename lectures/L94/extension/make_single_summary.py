from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree as E
import posixpath,hashlib,json
from pptx import Presentation
B=Path(__file__).resolve().parent.parent;src=B/'deliverables/Neuroethology_Lecture94_FA2026_Expanded71_Compatible.pptx';out=B/'deliverables/Neuroethology_Lecture94_FA2026_70slides_SingleSummary.pptx'
P='http://schemas.openxmlformats.org/presentationml/2006/main';R='http://schemas.openxmlformats.org/officeDocument/2006/relationships';A='http://schemas.openxmlformats.org/drawingml/2006/main'
def rels(part):
 p,n=posixpath.split(part);return posixpath.join(p,'_rels',n+'.rels')
def resolve(part,target):return posixpath.normpath(posixpath.join(posixpath.dirname(part),target))if not target.startswith('/')else target.lstrip('/')
with ZipFile(src)as z:
 pres=E.fromstring(z.read('ppt/presentation.xml'));pr=E.fromstring(z.read('ppt/_rels/presentation.xml.rels'));lst=pres.find('{'+P+'}sldIdLst');s=lst[45];rid=s.get('{'+R+'}id');relation=next(r for r in pr if r.get('Id')==rid);part=resolve('ppt/presentation.xml',relation.get('Target'));removed={part,rels(part)}
 for r in E.fromstring(z.read(rels(part))):
  if r.get('Type').endswith('/notesSlide'):
   n=resolve(part,r.get('Target'));removed.update([n,rels(n)])
 lst.remove(s);pr.remove(relation)
 edits={};mapping={};seen={}
 for n in z.namelist():
  if not n.startswith('ppt/media/'):continue
  h=hashlib.sha256(z.read(n)).hexdigest()
  if h in seen:mapping[n]=seen[h];removed.add(n)
  else:seen[h]=n
 for n in z.namelist():
  if n in removed or not n.endswith('.rels') or n=='ppt/_rels/presentation.xml.rels':continue
  root=E.fromstring(z.read(n));parent=posixpath.dirname(posixpath.dirname(n));changed=False
  for r in root:
   if r.get('TargetMode')=='External':continue
   full=posixpath.normpath(posixpath.join(parent,r.get('Target')))
   if full in mapping:r.set('Target',posixpath.relpath(mapping[full],parent));changed=True
  if changed:edits[n]=E.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
 for i,sid in enumerate(lst,1):
  rel=next(r for r in pr if r.get('Id')==sid.get('{'+R+'}id'));n=resolve('ppt/presentation.xml',rel.get('Target'))
  if i<46:continue
  root=E.fromstring(z.read(n))
  for sp in root.iter('{'+P+'}sp'):
   nv=sp.find('.//{'+P+'}cNvPr')
   if nv is not None and nv.get('name')=='Slide number':
    texts=list(sp.iter('{'+A+'}t'));assert len(texts)==1;texts[0].text=str(i)
  edits[n]=E.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
 ct=E.fromstring(z.read('[Content_Types].xml'))
 for x in list(ct):
  if x.get('PartName','').lstrip('/')in removed:ct.remove(x)
 app=E.fromstring(z.read('docProps/app.xml'))
 for x in app.iter():
  if x.tag.endswith('}Slides'):x.text='70'
 for n,r in [('ppt/presentation.xml',pres),('ppt/_rels/presentation.xml.rels',pr),('[Content_Types].xml',ct),('docProps/app.xml',app)]:edits[n]=E.tostring(r,xml_declaration=True,encoding='UTF-8',standalone=True)
 with ZipFile(out,'w',ZIP_DEFLATED,compresslevel=9)as w:
  for n in z.namelist():
   if n not in removed:w.writestr(n,edits.get(n,z.read(n)))
with ZipFile(src)as z,ZipFile(out)as w:
 assert w.testzip() is None
 for n in w.namelist():
  if n.startswith('ppt/media/'):assert z.read(n)==w.read(n)
p=Presentation(out);assert len(p.slides)==70
keys=[i for i,sl in enumerate(p.slides,1)if any(sh.has_text_frame and sh.text.strip()=='Key takeaways'for sh in sl.shapes)];assert keys==[70],keys
assert sum(sh.shape_type==16 for sl in p.slides for sh in sl.shapes)==38
report={'slides':70,'key_takeaways_slides':keys,'movie_placements':38,'all_retained_media_bytes_unchanged':True,'duplicate_media_removed':mapping,'bytes':out.stat().st_size,'sha256':hashlib.sha256(out.read_bytes()).hexdigest()}
(B/'extension/single_summary_validation.json').write_text(json.dumps(report,indent=2)+'\n');print(json.dumps(report,indent=2))
