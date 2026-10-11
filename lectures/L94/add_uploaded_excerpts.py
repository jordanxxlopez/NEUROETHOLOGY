"""Add instructor-uploaded excerpts after the original recordings and EVE excerpts.
Keep the tracked 17-video deck intact; write the expanded downloadable deck separately.
"""
from pathlib import Path
import json,hashlib,zipfile
from collections import defaultdict
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
B=Path(__file__).resolve().parent
m=json.loads((B/'media_sources.json').read_text());src=B/'Neuroethology_Lecture94_FA2026.pptx';p=Presentation(src)
assert sum(s.shape_type==16 for slide in p.slides for s in slide.shapes)==17
original_text=[[s.text for s in slide.shapes if s.has_text_frame]for slide in p.slides]
movies=[(i,s.shape_id,s.element.xml)for i,slide in enumerate(p.slides)for s in slide.shapes if s.shape_type==16]
groups=defaultdict(list)
for c in m['additional_uploaded_excerpts']:groups[c['slide']].append(c)
for number,clips in groups.items():
 slide=p.slides[number-1];existing=any(s.shape_type==16 for s in slide.shapes)
 pics=[s for s in slide.shapes if s.shape_type==13 and not s.name.startswith(('Unedited article-recording frame:','Unedited EVE'))]
 caps=[s for s in slide.shapes if s.name.startswith('Figure caption')]
 assert len(pics)==len(caps)
 for j,(pic,cap)in enumerate(zip(pics,caps)):
  ratio=pic.width/pic.height
  if existing:
   cell=2.75;left=7.01;top=1.35+j*1.4;limit=1.0 if len(pics)>1 else 2.2
  else:
   cell=5.7 if len(pics)==1 else 2.75;left=7.01+j*(2.95 if len(pics)>1 else 0);top=1.35;limit=2.5
  w=min(cell,limit*ratio);h=w/ratio
  pic.left=Inches(left+(cell-w)/2);pic.top=Inches(top);pic.width=Inches(w);pic.height=Inches(h)
  cap.left=Inches(left);cap.width=Inches(cell);cap.top=Inches(top+h+.05);cap.height=Inches(.35 if existing else .43)
  for para in cap.text_frame.paragraphs:
   for r in para.runs:r.font.size=Pt(8 if existing or len(pics)>1 else 9)
 for j,c in enumerate(clips):
  cell=2.75 if existing or len(clips)>1 else 5.7;left=9.96 if existing else 7.01+j*(2.95 if len(clips)>1 else 0);top=1.35 if existing else 4.23
  w=min(cell,2.35*c['display_aspect']);h=w/c['display_aspect'];x=left+(cell-w)/2
  f=B/c['path'];poster=B/c['poster'];assert hashlib.sha256(f.read_bytes()).hexdigest()==c['sha256']
  fallback=slide.shapes.add_picture(str(poster),Inches(x),Inches(top),Inches(w),Inches(h));fallback.name='Unedited uploaded recording source frame'
  movie=slide.shapes.add_movie(str(f),Inches(x),Inches(top),Inches(w),Inches(h),poster_frame_image=str(poster),mime_type='video/mp4');movie.name=c['caption']+' (click to play)'
  box=slide.shapes.add_textbox(Inches(left),Inches(top+h+.08),Inches(cell),Inches(.7));box.name='Uploaded excerpt credit';tf=box.text_frame;tf.word_wrap=True;tf.margin_top=tf.margin_bottom=tf.margin_left=tf.margin_right=0
  r=tf.paragraphs[0].add_run();r.text='Click to play. '+c['caption'];r.font.name='Arial';r.font.size=Pt(9);r.font.color.rgb=RGBColor(0,0,0)
  notes=slide.notes_slide.notes_text_frame
  for t in ['', 'Instructor-requested excerpt: '+c['caption'],'• '+c['teaching_note'],'Original filename: '+c['original_filename'],'Source interval: '+str(c['source_start_s'])+'–'+str(c['source_end_s'])+' seconds.','Publication provenance: '+c['publication_status'],'Credit: '+c['credit'],'License: '+c['license'],'Editing: '+c['conversion']]:notes.add_paragraph().text=t
  if c['doi']:notes.add_paragraph().text='Article DOI: https://doi.org/'+c['doi']+'; original supplement: '+c['original_media_url']
  else:notes.add_paragraph().text='This instructor-supplied contextual footage is not represented as an independently verified article supplement. The cited primary articles support the lecture’s scientific findings.'
  if number==23:notes.add_paragraph().text='Source-access update: the previously unavailable Roberts Video S1 has now been supplied by the instructor and its exact title was verified against the uploaded paper.'
for i,slide in enumerate(p.slides):assert [s.text for s in slide.shapes if s.has_text_frame][:len(original_text[i])]==original_text[i]
for i,shape_id,xml in movies:assert next(s for s in p.slides[i].shapes if s.shape_id==shape_id).element.xml==xml
out=B/'deliverables';out.mkdir(exist_ok=True);deck=out/'Neuroethology_Lecture94_FA2026.pptx';p.save(deck)
with zipfile.ZipFile(src)as old,zipfile.ZipFile(deck)as new:
 for n in old.namelist():
  if n.startswith('ppt/media/'):assert old.read(n)==new.read(n)
 hashes={hashlib.sha256(new.read(n)).hexdigest()for n in new.namelist()if n.endswith('.mp4')}
 assert hashes=={c['sha256']for c in m['embedded']+m['eve_excerpts']+m['additional_uploaded_excerpts']}
 for c in m['additional_uploaded_excerpts']:
  assert b'<p:timing>'in new.read(f"ppt/slides/slide{c['slide']}.xml")
  assert b'/video'in new.read(f"ppt/slides/_rels/slide{c['slide']}.xml.rels")
assert len(p.slides)==46
print('Expanded deck:',deck,'bytes:',deck.stat().st_size,'25 embedded videos; original text and 17 prior movies preserved.')
