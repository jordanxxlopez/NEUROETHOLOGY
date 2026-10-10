"""Embed the authors' unedited recording after running the repository builder.
The MP4 is a format conversion of the cited lab repository's original GIF.
Run: python lectures/L79/embed_video.py --poster /path/to/original/video_poster.png
"""
from pathlib import Path
import argparse,json,hashlib,zipfile
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
D=Path(__file__).resolve().parent
p=argparse.ArgumentParser();p.add_argument('--poster',default=str(D/'media/video_poster.png'));a=p.parse_args()
path=D/'Neuroethology_Lecture79_FA2026.pptx';r=Presentation(path);s=r.slides[1]
# Retain an intact article figure behind the media as a PDF/viewer fallback.
f=next(x for x in s.shapes if x.shape_type==MSO_SHAPE_TYPE.PICTURE)
f.top=Inches(1.4);f.height=Inches(5.6*1080/1370);f.width=Inches(3.994);f.left=Inches(7.823)
for shape in list(s.shapes):
 if shape.has_text_frame and shape.text.startswith('Givon et al. (2022), Fig.'):
  shape._element.getparent().remove(shape._element)
movie=s.shapes.add_movie(str(D/'media/Givon_author_side_view.mp4'),Inches(7.02),Inches(1.4),Inches(5.6),Inches(5.6*1080/1370),poster_frame_image=a.poster,mime_type='video/mp4')
b=s.shapes.add_textbox(Inches(7.02),Inches(5.95),Inches(5.6),Inches(.38));q=b.text_frame.paragraphs[0];q.text='Givon et al. (2022), Fig. 1(A–E). Static vehicle setup. Video: authors’ vehicle repository.';q.font.name='Arial';q.font.size=Pt(9);q.font.color.rgb=RGBColor(0,0,0)
s.notes_slide.notes_text_frame.text=s.notes_slide.notes_text_frame.text.split('Figure sources:')[0]
s.notes_slide.notes_text_frame.text+='\n\nVideo source: Givon et al. (2022), authors’ cited FishOperatedVehicle repository. Original side-view GIF: https://github.com/RonenSegevLab/FishOperatedVehicle/blob/main/videos/movie%20side%20view.gif . Runtime about 13 seconds. The original recording was converted to H.264 MP4 without trimming, recoloring, adding graphics or changing the scene. Click the recording during PowerPoint Slide Show to play it. The original Fig. 1(A–E) remains underneath as a static fallback for PDF export and viewers that omit video. The media poster is an original recording frame.\n'
r.save(path)
with zipfile.ZipFile(path) as z:
 media=[n for n in z.namelist() if n.endswith('.mp4')];assert len(media)==1;assert hashlib.sha256(z.read(media[0])).digest()==hashlib.sha256((D/'media/Givon_author_side_view.mp4').read_bytes()).digest();assert 'videoFile' in z.read('ppt/slides/slide2.xml').decode()
print('Embedded MP4 verified in slide 2:',media[0])
