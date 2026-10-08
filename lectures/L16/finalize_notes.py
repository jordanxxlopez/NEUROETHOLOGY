"""Retain the final slide teaching transcript and References heading."""
import json
from pathlib import Path
from pptx import Presentation
p=Path(__file__).resolve().parent
s=json.loads((p/'lecture.json').read_text());f=p/'Neuroethology_Lecture16_FA2026.pptx';r=Presentation(f);tk=s['takeaways'];r.slides[-1].notes_slide.notes_text_frame.text='Teaching transcript:\n'+'\n\n'.join('• '+x for x in tk['transcript'])+'\n\nReferences:\n'+'\n'.join(tk['refs']);r.save(f)
