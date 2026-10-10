"""Embed original article recordings in the repository-built PPTX."""
from pathlib import Path
import json,hashlib,zipfile,subprocess,sys
from PIL import Image
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
B=Path(__file__).resolve().parent
P=B/'Neuroethology_Lecture81_FA2026.pptx'
M=json.loads((B/'media_sources.json').read_text())
prs=Presentation(P)
for m in M['embedded']:
 video=B/m['path'];poster=B/m['poster']
 assert hashlib.sha256(video.read_bytes()).hexdigest()==m['sha256']
 s=prs.slides[m['slide']-1];pics=[sh for sh in s.shapes if sh.shape_type==13]
 assert len(pics)==1
 photo=pics[0];ratio=photo.width/photo.height
 w=min(5.8,2.2*ratio);h=w/ratio
 photo.width=Inches(w);photo.height=Inches(h);photo.left=Inches(6.983+(5.8-w)/2);photo.top=Inches(1.38)
 for sh in s.shapes:
  if sh.has_text_frame and ', Fig. ' in sh.text:
   sh.top=Inches(1.48+h);sh.height=Inches(.45)
 im=Image.open(poster);ratio=im.width/im.height;vh=2.30;vw=vh*ratio
 s.shapes.add_picture(str(poster),Inches(6.983+(5.8-vw)/2),Inches(4.10),Inches(vw),Inches(vh))
 movie=s.shapes.add_movie(str(video),Inches(6.983+(5.8-vw)/2),Inches(4.10),Inches(vw),Inches(vh),poster_frame_image=str(poster),mime_type='video/mp4')
 movie.name=m['caption']+' (embedded; click to play)'
 cap=s.shapes.add_textbox(Inches(6.983),Inches(6.48),Inches(5.8),Inches(.34))
 tf=cap.text_frame;tf.word_wrap=True;tf.margin_top=tf.margin_bottom=tf.margin_left=tf.margin_right=0
 r=tf.paragraphs[0].add_run();r.text=m['caption'];r.font.name='Arial';r.font.size=Pt(9);r.font.color.rgb=RGBColor(0,0,0)
 s.notes_slide.notes_text_frame.text+='\n\nArticle-provided recording:\n'+m['caption']+'\nArticle DOI: '+m['doi']+'\nOriginal media URL: '+m['original_media_url']+'\nCredit: '+m['credit']+'\nLicense: '+m['license']+'\n'+m['teaching_note']+'\n'+m['conversion']+'\nPlayback: embedded H.264/AAC MP4; click in PowerPoint Slide Show. The credited static source figure remains on the slide. The poster is an unedited original recording frame. Media decode verified; interactive PowerPoint playback was not available in this environment.'
spec=json.loads((B/'lecture.json').read_text())
prs.slides[-1].notes_slide.notes_text_frame.text='Teaching transcript:\n'+'\n'.join('• '+t for t in spec['takeaways']['transcript'])+'\n\nReferences:\n'+'\n\n'.join(spec['takeaways']['refs'])
prs.save(P)
with zipfile.ZipFile(P) as z:
 paths=[n for n in z.namelist() if n.endswith('.mp4')]
 assert len(paths)==2
 hashes={hashlib.sha256(z.read(n)).hexdigest() for n in paths}
 for m in M['embedded']:
  assert m['sha256'] in hashes
  rel=z.read(f"ppt/slides/_rels/slide{m['slide']}.xml.rels").decode()
  assert '/video' in rel and '/media' in rel
print('Verified two embedded original recordings, byte hashes, and media relationships')
subprocess.run([sys.executable,str(B.parents[1]/'tools/check_lecture.py'),str(P),'--lecture','81'],check=True)
