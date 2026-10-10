"""Use the repository builder; position the unusually long exact title safely."""
from pathlib import Path
import sys,runpy
from pptx.util import Inches
R=Path(__file__).resolve().parents[2];sys.path.insert(0,str(R/'tools'))
import build_lecture
original=build_lecture.Deck.title_slide
def title_slide(self):
 original(self)
 for sh in self.prs.slides[0].shapes:
  if not sh.has_text_frame:continue
  if sh.name=='Lecture title':sh.top=Inches(1.85)
  elif sh.text.startswith('NEUR 411'):sh.top=Inches(.65)
  elif sh.text=='Lecture 62':sh.top=Inches(1.35)
  elif sh.text=='Date TBD':sh.top=Inches(6.05)
  elif sh.text.startswith('Jordan Lopez'):sh.top=Inches(6.65)
build_lecture.Deck.title_slide=title_slide
sys.argv=['build_lecture.py',str(Path(__file__).with_name('lecture.json'))]
build_lecture.main()
