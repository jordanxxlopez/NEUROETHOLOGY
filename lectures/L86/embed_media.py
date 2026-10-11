"""Embed original article supplements with native proportions and verified provenance."""
from pathlib import Path
import json,subprocess,hashlib,zipfile
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
D=Path(__file__).resolve().parent
records=json.loads((D/'media_sources.json').read_text())
assign={
('dlc','1'):('Pose separates body configuration from location','Tri-mouse pose predictions.'),
('dlc','2'):('Parenting creates overlap between unlike bodies','Parenting mouse and pup predictions.'),
('dlc','3'):('Appearance helps preserve identity during social contact','Marmoset pose predictions.'),
('dlc','4'):('Schooling fish challenge a two-dimensional tracker','Schooling fish pose predictions.'),
('dlc','5'):('Silhouette tracking fails differently from pose tracking','idtracker.ai on mice and marmosets.'),
('sleap','1'):('A shared tracking system needs different body definitions','Uncorrected predictions across species.'),
('sleap','2'):('Human labels define what the computer learns to measure','Original labeling-session screencast.'),
('sleap','3'):('Grouping can precede or follow animal localization','Original authors’ pose-processing explanation.'),
('sleap','4'):('An identity error can persist after animals separate','Temporal versus appearance identity tracking.'),
('sleap','5'):('Social geometry specifies an intervention trigger','Approach-triggered ovipositor extrusion.')}
p=D/'Neuroethology_Lecture86_FA2026.pptx';prs=Presentation(p)
before=[[x.text for x in s.shapes if x.has_text_frame] for s in prs.slides]
# Preserve the long exact title while keeping date and instructor inside the canvas.
for shape in prs.slides[0].shapes:
 if shape.name=='Lecture title':shape.top=Inches(1.85)
 elif shape.has_text_frame and shape.text=='Lecture 86':shape.top=Inches(1.25)
 elif shape.has_text_frame and shape.text=='Date TBD':shape.top=Inches(6.02)
 elif shape.has_text_frame and 'Jordan Lopez' in shape.text:
  shape.top=Inches(6.7);shape.height=Inches(.55)
  for para in shape.text_frame.paragraphs:
   for run in para.runs:run.font.size=Pt(10)

for r in records:
 title,desc=assign[(r['key'],r['id'])];idx=next(i for i,s in enumerate(prs.slides) if any(x.has_text_frame and x.text==title for x in s.shapes));slide=prs.slides[idx]
 author='Lauer et al.' if r['key']=='dlc' else 'Pereira et al.';doi='10.1038/s41592-022-01443-0' if r['key']=='dlc' else '10.1038/s41592-022-01426-1'
 caption=f"{author} (2022), Supplementary Video {r['id']}. {desc}"
 figure=next(x for x in slide.shapes if x.shape_type==MSO_SHAPE_TYPE.PICTURE)
 cap=next(x for x in slide.shapes if x.has_text_frame and any(w in x.text for w in [', Fig.',', Extended Data Fig.']))
 # Source figure and movie occupy separate rectangles, never distort either.
 fx,fy,fw,fh=7.08,1.35,5.65,2.18
 ratio=figure.width/figure.height;nw,nh=min(fw,fh*ratio),min(fh,fw/ratio)
 figure.left,figure.top,figure.width,figure.height=map(Inches,(fx+(fw-nw)/2,fy+(fh-nh)/2,nw,nh))
 cap.left,cap.top,cap.width,cap.height=map(Inches,(7.08,3.57,5.65,.56))
 for paragraph in cap.text_frame.paragraphs:
  for run in paragraph.runs:run.font.size=Pt(9.5)
 v=next(x for x in r['probe']['streams'] if x['codec_type']=='video');vr=v['width']/v['height']
 x,y,w,h=7.08,4.18,5.65,2.32;vw,vh=min(w,h*vr),min(h,w/vr);x+=(w-vw)/2;y+=(h-vh)/2
 # Static fallback is an unchanged frame from the same movie.
 slide.shapes.add_picture(str(D/r['poster']),Inches(x),Inches(y),width=Inches(vw),height=Inches(vh))
 slide.shapes.add_movie(str(D/r['file']),Inches(x),Inches(y),Inches(vw),Inches(vh),poster_frame_image=str(D/r['poster']),mime_type='video/mp4')
 slide._element.findall('.//{http://schemas.openxmlformats.org/presentationml/2006/main}video')[-1].set('fullScrn','1')
 txt=slide.shapes.add_textbox(*map(Inches,(7.08,6.54,5.65,.33)));tf=txt.text_frame
 tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 q=tf.paragraphs[0];q.text=caption;q.font.name='Arial';q.font.size=Pt(8.5);q.font.color.rgb=RGBColor(0,0,0)
 r.update(slide=idx+1,caption=caption,doi=doi,credit=author+' and the original study authors; Nature Methods',license='CC BY 4.0, original article supplementary material',bounds_inches=[x,y,vw,vh],aspect_ratio=vr,full_screen_playback_requested=True)
 slide.notes_slide.notes_text_frame.text+='\n\nArticle-provided recording:\n'+caption+'\nArticle DOI: https://doi.org/'+doi+'\nOriginal media: '+r['original_url']+'\nOriginal filename: '+r['original_filename']+'\nCredit: '+r['credit']+'\nLicense: '+r['license']+'\n'+r['conversion']+'\nClick the recording in PowerPoint Slide Show for full-screen playback. The separate figure remains the credited static evidence. Embedded relationships and full media decoding were verified; PowerPoint GUI playback was unavailable. Google Slides imports can omit embedded videos: upload the separately supplied MP4 to Drive and insert it as a video. The PDF preview is static.\n'
prs.save(p)
with zipfile.ZipFile(p) as z:
 blobs=[z.read(n) for n in z.namelist() if n.endswith('.mp4')];assert len(blobs)==10
 assert {hashlib.sha256(v).hexdigest() for v in blobs}=={r['sha256'] for r in records}
 for r in records:
  assert b'videoFile' in z.read(f"ppt/slides/slide{r['slide']}.xml")
  assert b'relationships/video' in z.read(f"ppt/slides/_rels/slide{r['slide']}.xml.rels")
after=Presentation(p)
for i,s in enumerate(after.slides):assert [x.text for x in s.shapes if x.has_text_frame][:len(before[i])]==before[i]
(D/'media_sources.json').write_text(json.dumps(records,indent=2,ensure_ascii=False)+'\n')
(D/'MEDIA.md').write_text('# Lecture 86 original article videos\n\nTen original journal supplements are embedded as native video objects. Figures and videos have separate positions. Scenes, proportions, timing and existing annotations are preserved. Original authors’ method-explanation video is supplied unchanged, not recreated.\n\nClick a movie in PowerPoint Slide Show for full-screen playback. GUI playback was not available to test. The PDF preview is static. Google Slides may omit imported embedded videos; upload these MP4s to Drive and insert them using Insert > Video.\n\n'+'\n\n'.join(f"- Slide {r['slide']}: {r['caption']} [Download MP4]({Path(r['file']).as_posix()}) · [Original supplement]({r['original_url']}) · [Article](https://doi.org/{r['doi']})" for r in sorted(records,key=lambda x:x['slide']))+'\n')
print('Ten video embeds verified; slide text unchanged; native aspect ratios preserved. Bytes:',p.stat().st_size)
