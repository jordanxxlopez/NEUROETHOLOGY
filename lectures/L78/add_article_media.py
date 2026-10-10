"""Add verified article-provided media after the required repository builder."""
from pathlib import Path
import json,subprocess,sys,hashlib,zipfile
from pptx import Presentation
from pptx.util import Inches,Pt
b=Path(__file__).resolve().parent
p=b/'Neuroethology_Lecture78_FA2026.pptx'
m=json.loads((b/'media_sources.json').read_text())['embedded'][0]
video=b/m['path'];poster=b/'media/elephant_s1_poster.png'
assert hashlib.sha256(video.read_bytes()).hexdigest()==m['sha256']
if not poster.exists():subprocess.run(['ffmpeg','-v','error','-ss','2','-i',str(video),'-frames:v','1',str(poster)],check=True)
prs=Presentation(p);s=prs.slides[m['slide']-1]
pics=[sh for sh in s.shapes if sh.shape_type==13];assert len(pics)==1
photo=pics[0];ratio=photo.width/photo.height;photo.height=Inches(2.25);photo.width=int(photo.height*ratio);photo.left=Inches(6.92)+(Inches(5.8)-photo.width)//2;photo.top=Inches(1.38)
for sh in s.shapes:
 if sh.has_text_frame and 'Fig. 5A' in sh.text:sh.top=Inches(3.7);sh.height=Inches(.4)
movie=s.shapes.add_movie(str(video),Inches(7.84),Inches(4.15),Inches(3.96),Inches(2.2275),poster_frame_image=str(poster),mime_type='video/mp4')
movie.name='Cordoni 2025 Video S1 (embedded; click to play)'
cap=s.shapes.add_textbox(Inches(6.92),Inches(6.48),Inches(5.8),Inches(.35));tf=cap.text_frame;tf.word_wrap=True;tf.margin_top=tf.margin_bottom=0;tf.margin_left=tf.margin_right=0
r=tf.paragraphs[0].add_run();r.text=m['caption'];r.font.name='Arial';r.font.size=Pt(9)
s.notes_slide.notes_text_frame.text+='\n\nArticle-provided recording:\n'+m['caption']+'\nArticle DOI: '+m['doi']+'\nOriginal media URL: '+m['original_media_url']+'\nCredit: '+m['credit']+'\nLicense: '+m['license']+'\n'+m['teaching_note']+'\nPlayback: embedded H.264 MP4; click the video in PowerPoint Slide Show. Original published static figure retained. Poster is an unedited frame at 2 seconds. Decode verified; interactive PowerPoint playback was not available in this environment.'
prs.save(p)
with zipfile.ZipFile(p) as z:
 media=[n for n in z.namelist() if n.endswith('.mp4')];assert len(media)==1
 assert hashlib.sha256(z.read(media[0])).hexdigest()==m['sha256']
 xml=z.read('ppt/slides/_rels/slide39.xml.rels').decode();assert '/video' in xml and '/media' in xml
print('Verified embedded article video and slide39 media relationships')
subprocess.run([sys.executable,str(b.parents[1]/'tools/check_lecture.py'),str(p),'--lecture','78'],check=True)
