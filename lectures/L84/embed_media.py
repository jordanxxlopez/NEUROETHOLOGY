from pathlib import Path
import subprocess,json,hashlib,zipfile
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
D=Path(__file__).resolve().parent;Q=Path('/workspace/scratch/L84');M=D/'media';M.mkdir(exist_ok=True)
records=[
('A nose swipe changes where software registers a choice','pig_videos/S2.mp4','Ao2022_S2.mp4','Ao et al. (2022), Supplementary Video S2. Touchscreen response shaping.','10.3390/biomedicines10102612','https://mdpi-res.com/d_attachment/biomedicines/biomedicines-10-02612/article_deploy/biomedicines-10-02612-s001.zip',10),
('Correction trials support color discrimination','pig_videos/S3.mp4','Ao2022_S3.mp4','Ao et al. (2022), Supplementary Video S3. Pig color discrimination.','10.3390/biomedicines10102612','https://mdpi-res.com/d_attachment/biomedicines/biomedicines-10-02612/article_deploy/biomedicines-10-02612-s001.zip',15),
('Smaller dots train more precise screen contact','primate_video/Video_S1.mp4','Martin2022_VideoS1.mp4','Martin et al. (2022), Supplementary Video S1. Macaque touchscreen tasks.','10.3390/ani12131660','https://mdpi-res.com/d_attachment/animals/animals-12-01660/article_deploy/animals-12-01660-s001.zip',30),
('Rats move a vehicle by completing an electrical circuit','rat_movie2.mp4','Crawford2020_mmc2.mp4','Crawford et al. (2020), article video: mmc2.mp4. Forward driving.','10.1016/j.bbr.2019.112309','https://ars.els-cdn.com/content/image/1-s2.0-S0166432819311763-mmc2.mp4',4),
('Changed start positions recruit corrective steering','rat_movie1.mp4','Crawford2020_mmc1.mp4','Crawford et al. (2020), article video: mmc1.mp4. Changed starts and goals.','10.1016/j.bbr.2019.112309','https://ars.els-cdn.com/content/image/1-s2.0-S0166432819311763-mmc1.mp4',22)]
p=D/'Neuroethology_Lecture84_FA2026.pptx';prs=Presentation(p);before=[[x.text for x in s.shapes if x.has_text_frame] for s in prs.slides];out=[]
for title,src,name,caption,doi,url,second in records:
 index=next(i for i,s in enumerate(prs.slides) if any(x.has_text_frame and x.text==title for x in s.shapes));slide=prs.slides[index];movie=M/name
 if not movie.exists():
  subprocess.run(['ffmpeg','-v','error','-i',str(Q/src),'-c:v','libx264','-preset','fast','-crf','24','-pix_fmt','yuv420p','-c:a','copy','-movflags','+faststart','-y',str(movie)],check=True)
 probe=json.loads(subprocess.check_output(['ffprobe','-v','error','-select_streams','v:0','-show_entries','stream=width,height,sample_aspect_ratio','-of','json',str(movie)]))['streams'][0];ratio=probe['width']/probe['height']
 poster=M/(movie.stem+'_poster.jpg');subprocess.run(['ffmpeg','-v','error','-ss',str(second),'-i',str(movie),'-frames:v','1','-y',str(poster)],check=True)
 fig=next(x for x in slide.shapes if x.shape_type==MSO_SHAPE_TYPE.PICTURE);cap=next(x for x in slide.shapes if x.has_text_frame and ', Fig.' in x.text)
 def fit(shape,box):
  x,y,w,h=box;r=shape.width/shape.height;nw,nh=min(w,h*r),min(h,w/r);shape.left,shape.top,shape.width,shape.height=map(Inches,(x+(w-nw)/2,y+(h-nh)/2,nw,nh))
 if ratio<1:
  fit(fig,(7.08,1.4,2.75,4.75));cap.left,cap.top,cap.width,cap.height=map(Inches,(7.08,6.18,2.75,.56));x,y,w,h=(10.02,1.45,2.48,4.85);creditbox=(10.02,6.18,2.48,.56)
 else:
  fit(fig,(7.08,1.35,5.44,2.65));cap.left,cap.top,cap.width,cap.height=map(Inches,(7.08,4.04,5.44,.49));x,y,w,h=(7.08,4.6,5.44,1.87);creditbox=(7.08,6.54,5.44,.28)
 for para in cap.text_frame.paragraphs:
  for run in para.runs:run.font.size=Pt(9)
 vw,vh=min(w,h*ratio),min(h,w/ratio);x+=(w-vw)/2;y+=(h-vh)/2
 slide.shapes.add_picture(str(poster),Inches(x),Inches(y),width=Inches(vw),height=Inches(vh))
 slide.shapes.add_movie(str(movie),Inches(x),Inches(y),Inches(vw),Inches(vh),poster_frame_image=str(poster),mime_type='video/mp4')
 slide._element.findall('.//{http://schemas.openxmlformats.org/presentationml/2006/main}video')[-1].set('fullScrn','1')
 txt=slide.shapes.add_textbox(*map(Inches,creditbox));txt.text_frame.margin_left=txt.text_frame.margin_right=0;txt.text_frame.margin_top=txt.text_frame.margin_bottom=0
 para=txt.text_frame.paragraphs[0];para.text=caption;para.font.name='Arial';para.font.size=Pt(8.5);para.font.color.rgb=RGBColor(0,0,0)
 license='CC BY 4.0, article supplementary material' if 'mdpi' in url else 'Copyright Elsevier 2019; no separate reuse license stated in supplied article.'
 context='The original study recording retains its original sequence, audio, timing and source annotations. Click during PowerPoint Slide Show for full-screen playback. Its poster is an unedited frame. Google Slides imports can omit embedded video; upload the supplied MP4 to Drive and insert it as a video. PowerPoint GUI playback was unavailable for testing.'
 slide.notes_slide.notes_text_frame.text+='\n\nArticle-provided recording:\n'+caption+'\n'+context+'\nDOI: https://doi.org/'+doi+'\nOriginal media: '+url+'\nCredit: original study authors and publisher.\nLicense: '+license+'\n'
 subprocess.run(['ffmpeg','-v','error','-i',str(movie),'-f','null','-'],check=True)
 out.append({'slide':index+1,'file':'media/'+name,'caption':caption,'doi':doi,'original_url':url,'original_filename':Path(src).name,'credit':'Original study authors and publisher','license':license,'conversion':'H.264 delivery encoding at original pixel dimensions; original sequence, timing, frame rate and audio retained; no cropping, recoloring, annotations or speed changes.','sha256':hashlib.sha256(movie.read_bytes()).hexdigest(),'bounds_inches':[x,y,vw,vh],'aspect_ratio':ratio,'duration_seconds':float(subprocess.check_output(['ffprobe','-v','error','-show_entries','format=duration','-of','default=nw=1:nk=1',str(movie)])),'complete_decode_verified':True})
prs.save(p)
with zipfile.ZipFile(p) as z:
 payloads=[z.read(n) for n in z.namelist() if n.endswith('.mp4')];assert len(payloads)==5;assert {hashlib.sha256(b).hexdigest() for b in payloads}=={r['sha256'] for r in out}
 for r in out:
  xml=z.read(f"ppt/slides/slide{r['slide']}.xml");rels=z.read(f"ppt/slides/_rels/slide{r['slide']}.xml.rels");assert b'videoFile' in xml and b'relationships/video' in rels
new=Presentation(p)
for i,s in enumerate(new.slides):assert [x.text for x in s.shapes if x.has_text_frame][:len(before[i])]==before[i]
(D/'media_sources.json').write_text(json.dumps(out,indent=2))
print('Embedded five original study videos; all slide text retained; embedded hashes and relationships verified.')
