"""Embed 20 verified article recordings, retaining scientific figures and native proportions."""
from pathlib import Path
import json,hashlib,subprocess
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
R=Path(__file__).resolve().parent;p=R/'Neuroethology_Lecture104A_FA2026.pptx'
rows=json.loads((R/'media_preparation.json').read_text());byname={Path(x['file']).stem:x for x in rows};prs=Presentation(p)
assert len(prs.slides)==46
assert not any(x.shape_type==MSO_SHAPE_TYPE.MEDIA for s in prs.slides for x in s.shapes),'Rebuild before embedding.'
entries={3:[('Arenas2017_Video1','Arenas et al. (2017), video 1. Planarian heat avoidance.')],4:[('Arenas2017_Video2','Arenas et al. (2017), video 2. Escape from an AITC-containing chamber.')],6:[('Khuong2019_MovieS1','Khuong et al. (2019), Movie S1. Adult fly thermal escape.')],7:[('Khuong2019_MovieS2','Khuong et al. (2019), Movie S2. Injury-induced thermal allodynia.')],10:[('Elife85760_Video1','Oikawa et al. (2023), Video 1. Control Goro calcium responses.')],11:[('Elife85760_Video2','Oikawa et al. (2023), Video 2. CCKLR-17D1-mutant Goro responses.')],12:[('Elife85760_Video3','Oikawa et al. (2023), Video 3. RNAi control.'),('Elife85760_Video4','Oikawa et al. (2023), Video 4. CCKLR-17D1 RNAi.')],13:[('Elife85760_Video5','Oikawa et al. (2023), Video 5. Receptor control.'),('Elife85760_Video6','Oikawa et al. (2023), Video 6. CCKLR-17D3 mutant.')],15:[('Worm2022_S3','Florman & Alkema (2022), S3 Video. Wild-type touch-evoked escape.')],16:[('Worm2022_S4','Florman & Alkema (2022), S4 Video. FLP-18 overexpression alters escape.')],17:[('Worm2022_S1','Florman & Alkema (2022), S1 Video. FLP-18 overexpression.'),('Worm2022_S2','Florman & Alkema (2022), S2 Video. Interneuron FLP-18 expression.')],30:[('Carrillo2019_VideoS1','Carrillo et al. (2019), Video S1. Observed shock and self-stimulation conditions.')],31:[('Carrillo2019_Fig2360','Carrillo et al. (2019), Figure360, Fig. 2. Authors explain population responses.')],32:[('Carrillo2019_Fig3360','Carrillo et al. (2019), Figure360, Fig. 3. Authors explain single-neuron selectivity.')],33:[('Carrillo2019_Fig4360','Carrillo et al. (2019), Figure360, Fig. 4. Authors explain behavioral responses.')],42:[('Bartal2011_MovieS1','Ben-Ami Bartal et al. (2011), Movie S1. Exploration of a trapped cagemate.')],43:[('Bartal2011_MovieS2','Ben-Ami Bartal et al. (2011), Movie S2. Opening the restrainer to release a cagemate.')]}
assert sum(map(len,entries.values()))==20
spec=json.loads((R/'lecture.json').read_text());refs=json.loads((R/'reference_verification.json').read_text());sources=[]
def fit(pic,x,y,w,h):
 ratio=pic.width/pic.height;fw=min(w,h*ratio);fh=fw/ratio;pic.left=Inches(x+(w-fw)/2);pic.top=Inches(y+(h-fh)/2);pic.width=Inches(fw);pic.height=Inches(fh)
def cap(sl,text,x,y,w,h,name='Video caption'):
 box=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));box.name=name;tf=box.text_frame;tf.word_wrap=True;tf.margin_top=tf.margin_bottom=tf.margin_left=tf.margin_right=0;tf.text=text
 for para in tf.paragraphs:
  for run in para.runs:run.font.name='Arial';run.font.size=Pt(9);run.font.color.rgb=RGBColor(0,0,0)
 return box
for n,items in entries.items():
 sl=prs.slides[n-1];before=[sh.text for sh in sl.shapes if sh.has_text_frame];pics=[sh for sh in sl.shapes if sh.shape_type==MSO_SHAPE_TYPE.PICTURE];captions=[sh for sh in sl.shapes if sh.name=='Figure caption']
 # The source figure remains distinct from its companion image and actual movie.
 if len(pics)>1:
  fit(pics[0],7.11,1.34,3.92,2.40);fit(pics[1],11.18,1.34,1.52,2.40)
  for j,(x,w) in enumerate([(7.11,3.92),(11.18,1.52)]):
   captions[j].left=Inches(x);captions[j].top=Inches(3.77);captions[j].width=Inches(w);captions[j].height=Inches(.61)
   for pa in captions[j].text_frame.paragraphs:
    for ru in pa.runs:ru.font.size=Pt(8)
 else:
  fit(pics[0],7.11,1.34,5.59,2.50);captions[0].left=Inches(7.11);captions[0].top=Inches(3.87);captions[0].width=Inches(5.59);captions[0].height=Inches(.48)
 count=len(items);slotw=(5.59-.18*(count-1))/count
 for j,(stem,caption) in enumerate(items):
  row=dict(byname[stem]);file=R/row['delivery_file'];post=R/row['poster'];probe=json.loads(subprocess.check_output(['ffprobe','-v','error','-show_streams','-of','json',str(file)]));v=next(x for x in probe['streams'] if x['codec_type']=='video');dar=v.get('display_aspect_ratio');ratio=float(dar.split(':')[0])/float(dar.split(':')[1]) if dar and dar!='0:1' else v['width']/v['height']
  maxh=2.10;w=min(slotw,maxh*ratio);h=w/ratio;x=7.11+j*(slotw+.18)+(slotw-w)/2;y=4.48+(maxh-h)/2
  fallback=sl.shapes.add_picture(str(post),Inches(x),Inches(y),width=Inches(w),height=Inches(h));fallback.name='Static video fallback'
  movie=sl.shapes.add_movie(str(file),Inches(x),Inches(y),Inches(w),Inches(h),poster_frame_image=str(post),mime_type='video/mp4');movie.name=stem
  for el in movie._element.xpath('.//p:cNvPr'):el.set('descr',caption+' Click to play the embedded original recording.')
  for el in sl._element.xpath('.//p:video'):el.set('fullScrn','1')
  cap(sl,caption,7.11+j*(slotw+.18),6.57,slotw,.38)
  doi=row['article_doi'];ref=next(x for x in refs.values() if x['doi'].lower()==doi.lower());credit=ref.get('authorString','Original study authors')
  license='CC BY 4.0' if stem.startswith(('Carrillo','Elife','Worm')) else 'CC BY-NC 4.0' if stem.startswith('Khuong') else 'No open license stated; credit original study authors'
  row.update({'slide':n,'caption':caption,'credit':credit,'license':license,'embedded':True,'native_aspect_ratio':ratio,'genre':row.get('genre','original experimental recording')});sources.append(row)
  sl.notes_slide.notes_text_frame.text+='\n\nEmbedded video: '+caption+'\nExact source identifier: '+row['original_identifier']+'\nArticle DOI: https://doi.org/'+doi+'\nOriginal media URL: '+row['download_url']+'\nCredit: '+credit+' License: '+license+'.\n'+row['conversion']+'. Poster is a source frame. Playback is full screen where the viewer supports it. Desktop PowerPoint GUI playback and Google Slides import remain untested. Google Slides may require separate Drive media insertion.\n'
 assert [sh.text for sh in sl.shapes if sh.has_text_frame][:len(before)]==before,f'Slide {n} teaching text changed'
prs.save(p);(R/'media_sources.json').write_text(json.dumps({'embedded_video_count':len(sources),'distinct_original_recordings':len({x['original_sha256'] for x in sources}),'experimental_recordings':17,'author_figure_presentations':3,'sources':sources},indent=2)+'\n');print('Embedded 20 recordings on',len(entries),'slides. PPTX bytes:',p.stat().st_size)
