"""Preserve the exact title in one editable paragraph and link note references."""
import json,re
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
base=Path(__file__).resolve().parent;path=base/'Neuroethology_Lecture51_FA2026.pptx';r=Presentation(path)
schedule=json.loads((base.parents[1]/'course/schedule.json').read_text());s=next(q for q in schedule['special_topics'] if q['n']==51);spec=json.loads((base/'lecture.json').read_text())
for sh in r.slides[0].shapes:
 if sh.name=='Lecture title':
  sh.top=Inches(2.5);sh.height=Inches(2.9);tf=sh.text_frame;tf.clear();tf.word_wrap=True;p=tf.paragraphs[0];p.text=s['title'];p.line_spacing=1.05;p.space_after=Pt(0)
  for run in p.runs:run.font.name='Arial';run.font.size=Pt(30);run.font.bold=True;run.font.color.rgb=RGBColor.from_string('FFFFFF')
 elif sh.has_text_frame and sh.text=='Date TBD':sh.top=Inches(5.8)
 elif sh.has_text_frame and sh.text.startswith(schedule['instructor']):sh.top=Inches(6.4)
r.slides[-1].notes_slide.notes_text_frame.text='\n'.join('• '+q['lead']+' '+q['text'] for q in spec['takeaways']['items'])+'\n\nReferences:\n'+'\n'.join(spec['takeaways']['refs'])
for sl in r.slides:
 for p in sl.notes_slide.notes_text_frame.paragraphs:
  text=p.text;p.clear()
  for part in re.split(r'(https://doi\.org/\S+)',text):
   if part:
    run=p.add_run();run.text=part;run.font.name='Arial';run.font.size=Pt(12)
    if part.startswith('https://doi.org/'):run.hyperlink.address=part
r.save(path)
