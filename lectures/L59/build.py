"""Build with the repository tool and keep the long-title instructor label inside the slide."""
from pathlib import Path
import subprocess,sys
from pptx import Presentation
from pptx.util import Inches
root=Path(__file__).resolve().parents[2]
spec=Path(__file__).with_name('lecture.json')
subprocess.run([sys.executable,str(root/'tools/build_lecture.py'),str(spec)],check=True)
p=spec.with_name('Neuroethology_Lecture59_FA2026.pptx');deck=Presentation(p)
for sh in deck.slides[0].shapes:
    if sh.has_text_frame and 'Professor of Neurobiology' in sh.text:
        sh.top=Inches(6.85)
        sh.height=Inches(.6)
deck.save(p)
subprocess.run([sys.executable,str(root/'tools/check_lecture.py'),str(p),'--lecture','59'],check=True)
