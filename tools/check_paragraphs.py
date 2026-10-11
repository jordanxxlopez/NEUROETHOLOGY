#!/usr/bin/env python3
"""Every slide except the first and last holds exactly three body paragraphs (scaffolding/RULES.md).

    python tools/check_paragraphs.py                 # all scaffolded/*.pptx
    python tools/check_paragraphs.py deck.pptx ...    # named decks

Prints one line per slide that breaks the rule. Exit code 1 when any slide is flagged.
"""
import glob
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
import scaffold_deck as sd  # noqa: E402
from pptx import Presentation  # noqa: E402


def body_paragraphs(slide):
    """Body paragraphs of one slide, including long text stranded in a caption-like box."""
    c = sd.classify(slide)
    if c["title"] is not None and c["title"].text_frame.text.strip().lower().startswith("key takeaways"):
        return None                                   # the takeaway grid is a card layout
    out = [p for b in c["body"] for p in b.text_frame.paragraphs if p.text.strip()]
    out += [p for b in c["caption"] if "Fig." not in b.text_frame.text and len(b.text_frame.text) > 60
            for p in b.text_frame.paragraphs if p.text.strip()]
    return [p for p in out if not is_caption_or_credit(p.text)]


def is_caption_or_credit(text):
    """Figure captions ('Author (year), Fig. N.') and photo credits are not body paragraphs."""
    t = text.strip()
    return bool(sd.CAPTION_START.search(t)) or t.startswith(("Photo:", "Credit:", "Left:", "Right:"))


def check(path):
    prs = Presentation(str(path))
    problems = []
    for k, slide in enumerate(prs.slides, 1):
        if k in (1, len(prs.slides)):
            continue
        paras = body_paragraphs(slide)
        if paras is None or len(paras) == 3:
            continue
        problems.append(f"{Path(path).name} slide {k}: {len(paras)} body paragraphs (expected 3)")
    return problems


if __name__ == "__main__":
    root = Path(__file__).resolve().parent.parent
    decks = sys.argv[1:] or sorted(glob.glob(str(root / "scaffolded" / "*.pptx")))
    bad = [line for d in decks for line in check(d)]
    print("\n".join(bad) if bad else f"all {len(decks)} decks: every middle slide has three body paragraphs")
    sys.exit(1 if bad else 0)
