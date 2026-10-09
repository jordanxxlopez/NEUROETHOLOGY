#!/usr/bin/env python3
"""Check a finished NEUR 411 deck against the house rules.

    python tools/check_lecture.py deck.pptx --lecture 11 [--content-slides 44]

Works on any .pptx, including decks not made by build_lecture.py.
Exit code 1 if any rule fails.
"""
import argparse
import json
import re
import sys
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent

sys.path.insert(0, str(Path(__file__).resolve().parent))
from style_rules import (MADE_IMAGE_WORDS, MIN_ARTICLE_FIGURE_SLIDES, MIN_COLOR_SHARE, MIN_IMAGE_SLIDES,  # noqa: E402
                         framing_problems, is_color_image, stats_problems)
from theme_colors import alignment_problems  # noqa: E402
from style_rules import EM_DASH, MAX_LIMITATION_SHARE, ends_with_limitation  # noqa: E402

STRUCTURE = [
    (r"^\s*(lecture )?(outline|agenda|overview|roadmap)\s*$", "no outline/agenda slide"),
    (r"\blearning (objectives|goals|outcomes)\b|^\s*objectives\s*$", "no objectives slide"),
    (r"\b(continued|cont\.)\s*$", "no 'continued' slide splits"),
]
MIN_TRANSCRIPT_WORDS = 60  # speaker-note teaching transcript per content slide
MIN_WORDS = 70         # per content slide, text only


def slide_text(slide, teaching_only=False):
    """All slide text; teaching_only skips figure captions and the citation footer."""
    out = []
    for sh in slide.shapes:
        if teaching_only and sh.name in ("Figure caption", "Citation", "Slide number"):
            continue
        if sh.has_text_frame:
            out.append(sh.text_frame.text)
        if sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    out.append(c.text)
    return "\n".join(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("deck")
    ap.add_argument("--lecture", type=int, required=True)
    ap.add_argument("--content-slides", type=int, default=44)
    ap.add_argument("--no-transcript", action="store_true",
                    help="skip the speaker-note transcript requirement (decks made before the rule)")
    a = ap.parse_args()

    sched = json.loads((ROOT / "course/schedule.json").read_text())
    meta = next(l for l in sched["lectures"] + sched.get("special_topics", []) if l["n"] == a.lecture)
    prs = Presentation(a.deck)
    slides = list(prs.slides)
    errs, warns = [], []

    want = a.content_slides + 2
    if len(slides) != want:
        errs.append(f"deck has {len(slides)} slides; need {want} (1 title + {a.content_slides} content + 1 key takeaways)")

    first = re.sub(r"\s+", " ", slide_text(slides[0]))
    if re.sub(r"\s+", " ", meta["title"]) not in first:
        errs.append(f"title slide does not contain the exact schedule title: {meta['title']!r}")
    # e.g. "September 18"; special topics carry "Date TBD"
    month_day = meta["date"].split(", ")[1] if ", " in meta["date"] else meta["date"]
    if month_day not in first:
        errs.append(f"title slide missing date {meta['date']!r}")

    if "key takeaways" not in slide_text(slides[-1]).lower():
        errs.append("last slide must be 'Key takeaways'")

    fig_slides = 0
    img_slides = 0
    color_slides = 0
    for i, s in enumerate(slides, 1):
        txt = slide_text(s)
        teach = slide_text(s, teaching_only=True)
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
        transcript = notes.split("References:")[0] if "References:" in notes else ""
        for line in txt.splitlines():
            for pat, why in STRUCTURE:
                if re.search(pat, line, re.I):
                    errs.append(f"slide {i}: {why}: {line.strip()[:80]!r}")
        for where, text in (("slide", teach), ("notes", transcript)):
            for why, hit in framing_problems(text):
                errs.append(f"slide {i} {where}: {why}: {hit!r}")
            for hit in stats_problems(text):
                errs.append(f"slide {i} {where}: statistics clutter {hit!r}; teach the finding, not test statistics")
        if i == 1:
            continue
        for line in notes.splitlines():
            if line.startswith(("Figure:", "Image:")) and MADE_IMAGE_WORDS.search(line):
                errs.append(f"slide {i}: image is self-made ({line[:70]!r}); only article figures or credited photos")
        if i == len(slides):
            continue
        words = len(teach.split())
        if words < MIN_WORDS:
            errs.append(f"slide {i}: only {words} words; content slides need >= {MIN_WORDS}")
        if not a.no_transcript and len(transcript.split()) < MIN_TRANSCRIPT_WORDS:
            errs.append(f"slide {i}: speaker notes need a teaching transcript (>= {MIN_TRANSCRIPT_WORDS} words) before 'References:'")
        pics = [sh for sh in s.shapes if sh.shape_type == 13 and sh.height > 914400 * 0.6]
        has_article = "Figure:" in notes
        has_credit_lines = has_article or "Image:" in notes
        if pics:
            img_slides += 1
            if any(is_color_image(sh.image.blob) for sh in pics):
                color_slides += 1
        if pics and (has_article or not has_credit_lines):
            fig_slides += 1
        if not re.search(r"\((19|20)\d{2}[a-z]?\)|\b(19|20)\d{2}[a-z]?\b", notes):
            errs.append(f"slide {i}: speaker notes need the full reference(s) with year")
        if "doi" not in notes.lower() and "http" not in notes.lower():
            warns.append(f"slide {i}: no DOI/URL in notes")

    n = a.content_slides
    need_img, need_art = round(MIN_IMAGE_SLIDES * n / 44), round(MIN_ARTICLE_FIGURE_SLIDES * n / 44)
    if img_slides < need_img:
        errs.append(f"only {img_slides} of {n} content slides have an image; nearly all must (>= {need_img})")
    if fig_slides < need_art:
        errs.append(f"only {fig_slides} content slides have article figures; need >= {need_art}")

    # caveat endings and em dashes (slide text only; captions, footers and notes excluded)
    caveats, dashes = [], []
    for i, sl in enumerate(slides[1:-1], start=2):
        frames = [sh for sh in sl.shapes if sh.has_text_frame
                  and sh.name not in ("Figure caption", "Citation", "Slide number")]
        if any(EM_DASH in sh.text_frame.text for sh in frames):
            dashes.append(i)
        body = [sh for sh in frames if sh.name == "Body"] or sorted(
            frames, key=lambda sh: len(sh.text_frame.text.split()), reverse=True)[:1]
        if body and ends_with_limitation([p.text for p in body[0].text_frame.paragraphs]):
            caveats.append(i)
    if dashes:
        errs.append(f"em dash in slide text on slides {dashes[:12]}; use a comma, semicolon, colon or new sentence")
    if len(caveats) > MAX_LIMITATION_SHARE * (len(slides) - 2):
        errs.append(f"{len(caveats)} slides end on a caveat paragraph ({caveats[:12]}); at most 1 in 5. The last "
                    "paragraph should teach; keep a limitation only when it changes what students should conclude")

    # one color per lecture: titles/bold = darker title-slide color, other text black
    errs.extend(alignment_problems(prs))

    if img_slides and color_slides < MIN_COLOR_SHARE * img_slides:
        warns.append(f"only {color_slides} of {img_slides} image slides have color figures; prefer colorful article figures "
                     "(micrographs, color maps, heat maps, color plots) — color only as published, never recolored")

    for w in warns:
        print("WARN ", w)
    for e in errs:
        print("FAIL ", e)
    print(f"{len(slides)} slides, {img_slides} image slides ({fig_slides} article-figure, {color_slides} in color), {len(errs)} failures, {len(warns)} warnings")
    sys.exit(1 if errs else 0)


if __name__ == "__main__":
    main()
