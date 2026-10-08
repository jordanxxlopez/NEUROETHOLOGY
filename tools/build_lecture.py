#!/usr/bin/env python3
"""Build a NEUR 411 lecture deck in the Lecture 8-10 format.

    python tools/build_lecture.py lectures/L11/lecture.json

The spec file (see lectures/_template/lecture.json) holds the content. Title,
date and course line are read from course/schedule.json so the lecture title
can never drift from the schedule. Output: 1 title slide + N content slides
(default 44) + 1 key-takeaways slide, 16:9, Arial, full references in notes.
After writing the deck the script runs tools/check_lecture.py on it.
"""
import json
import math
import re
import subprocess
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
FONT = "Arial"
W, H = 13.333, 7.5
M = 0.62                      # side margin (matches Lecture 10)
TITLE_Y, TITLE_H = 0.36, 0.62
BODY_Y = 1.35
FOOT_Y = 6.98                 # hairline above the citation footer
BODY_BOTTOM = 6.80
BODY_MAX_PT, BODY_MIN_PT = 18, 13
CAPTION_PT = 10
DEFAULT_CONTENT_SLIDES = 44


class SpecError(Exception):
    pass


# ---------------------------------------------------------------- helpers
def rgb(hex_):
    return RGBColor.from_string(hex_)


def inline_runs(text):
    """Split '**bold** and _italic_ text' into (chunk, bold, italic) runs."""
    out = []
    for part in re.split(r"(\*\*.+?\*\*|(?<![A-Za-z0-9])_[^_]+?_(?![A-Za-z0-9]))", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            out.append((part[2:-2], True, False))
        elif part.startswith("_") and part.endswith("_") and len(part) > 2:
            out.append((part[1:-1], False, True))
        else:
            out.append((part, False, False))
    return out


def plain(text):
    return "".join(c for c, _, _ in inline_runs(text))


def lines_needed(paras, width_in, pt):
    """Rough Arial line count: average glyph ≈ 0.5 em."""
    chars_per_line = max(1, int(width_in * 72 / (pt * 0.5)))
    total = 0
    for p in paras:
        words = plain(p).split()
        line, n = 0, 1
        for w in words:
            add = len(w) + (1 if line else 0)
            if line + add > chars_per_line:
                n, line = n + 1, len(w)
            else:
                line += add
        total += n
    return total


def fit_size(paras, width_in, height_in, max_pt=BODY_MAX_PT, min_pt=BODY_MIN_PT, where=""):
    for pt in range(max_pt, min_pt - 1, -1):
        n_lines = lines_needed(paras, width_in, pt)
        need = n_lines * pt * 1.18 / 72 + (len(paras) - 1) * pt * 0.6 / 72
        if need <= height_in * 0.97:
            return pt
    raise SpecError(
        f"{where}: text does not fit at {min_pt} pt in a {width_in:.1f}x{height_in:.1f} in box. "
        "Cut words, move a paragraph to another slide, or use the 'text' layout."
    )


def textbox(slide, x, y, w, h, name=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tb, tf


def write_paras(tf, paras, pt, color, bold_color=None, space_after=None, align=PP_ALIGN.LEFT):
    first = True
    for text in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(space_after if space_after is not None else pt * 0.6)
        p.line_spacing = 1.08
        for chunk, b, i in inline_runs(text):
            r = p.add_run()
            r.text = chunk
            f = r.font
            f.name, f.size, f.bold, f.italic = FONT, Pt(pt), b, i
            f.color.rgb = rgb(bold_color if (b and bold_color) else color)


def rect(slide, x, y, w, h, fill, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(0.75)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    style = s._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style")
    if style is not None:  # drop theme effect/line refs so no renderer adds a shadow
        s._element.remove(style)
    return s


def picture_contain(slide, path, x, y, w, h):
    """Place an image inside the box, preserving aspect ratio, top-centered."""
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    px = x + (w - pw) / 2
    slide.shapes.add_picture(str(path), Inches(px), Inches(y), Inches(pw), Inches(ph))
    return ph


def set_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


# ------------------------------------------------------------ slide parts
class Deck:
    def __init__(self, spec, spec_dir):
        self.spec = spec
        self.dir = spec_dir
        sched = json.loads((ROOT / "course/schedule.json").read_text())
        self.course = sched
        n = spec["lecture"]
        match = [l for l in sched["lectures"] if l["n"] == n]
        if not match:
            raise SpecError(f"Lecture {n} not in course/schedule.json")
        self.meta = match[0]
        themes = json.loads((ROOT / "course/themes.json").read_text())
        key = spec.get("theme")
        if key not in themes["palettes"]:
            raise SpecError(f"theme '{key}' not in course/themes.json palettes: {list(themes['palettes'])}")
        used_labels = {str(v) for v in themes["used"].values()}
        if themes["palettes"][key]["label"] in used_labels and not spec.get("allow_reused_theme"):
            raise SpecError(f"theme '{key}' was already used by another lecture; pick a new one")
        self.t = themes["palettes"][key]
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(H)
        self.blank = self.prs.slide_layouts[6]
        self.num = 0

    def path(self, p):
        q = (self.dir / p).resolve()
        if not q.exists():
            raise SpecError(f"figure not found: {q}")
        return q

    def new_slide(self):
        self.num += 1
        s = self.prs.slides.add_slide(self.blank)
        set_bg(s, self.t["bg"])
        return s

    def title(self, s, text):
        _, tf = textbox(s, M, TITLE_Y, W - 2 * M, TITLE_H, "Title")
        tf.vertical_anchor = MSO_ANCHOR.TOP
        write_paras(tf, [text], 30, self.t["heading"])
        for r in tf.paragraphs[0].runs:
            r.font.bold = True
        if len(text) > 62:
            raise SpecError(f"slide {self.num}: title over 62 characters will wrap: {text!r}")

    def footer(self, s, cite):
        rect(s, M, FOOT_Y, W - 2 * M, 0.01, self.t["rule"])  # hairline above the footer
        _, tf = textbox(s, M, FOOT_Y + 0.08, W - 2 * M - 0.6, 0.25, "Citation")
        write_paras(tf, [cite], 9, self.t["muted"], space_after=0)
        _, tf = textbox(s, W - M - 0.5, FOOT_Y + 0.08, 0.5, 0.25, "Slide number")
        write_paras(tf, [f"{self.num:02d}"], 9, self.t["muted"], space_after=0, align=PP_ALIGN.RIGHT)

    def notes(self, s, sd):
        parts = []
        parts += sd.get("refs", [])
        figs = [sd["figure"]] if "figure" in sd else sd.get("figures", [])
        for f in figs:
            parts.append("Figure: " + f["caption"])
            if f.get("source_url"):
                parts.append("Figure source: " + f["source_url"])
        if sd.get("notes"):
            parts.append(sd["notes"])
        s.notes_slide.notes_text_frame.text = "\n\n".join(parts)

    def caption(self, s, x, y, w, text):
        _, tf = textbox(s, x, y, w, 0.55, "Figure caption")
        write_paras(tf, [text], CAPTION_PT, self.t["muted"], space_after=0)

    def body(self, s, paras, x, y, w, h, where):
        pt = fit_size(paras, w, h, where=where)
        _, tf = textbox(s, x, y, w, h, "Body")
        write_paras(tf, paras, pt, self.t["text"], bold_color=self.t["heading"])
        return pt

    # ---------------------------------------------------------- layouts
    def title_slide(self):
        s = self.new_slide()
        t = self.t
        img = self.spec.get("title_image")
        panel_w = W / 2 if img else W
        rect(s, 0, 0, panel_w, H, t["title_bg"])
        x, w = 0.75, panel_w - 1.5
        _, tf = textbox(s, x, 0.9, w, 0.35)
        write_paras(tf, [self.course["course"]], 12, t["title_muted"], space_after=0)
        _, tf = textbox(s, x, 1.95, w, 0.45)
        write_paras(tf, [f"Lecture {self.meta['n']}"], 18, t["title_muted"], space_after=0)
        title = self.meta["title"]
        # break after the first colon for the two-line look; text is unchanged
        lines = [title[: title.index(":") + 1], title[title.index(":") + 1:].strip()] if ":" in title else [title]
        _, tf = textbox(s, x, 2.55, w, 2.1, "Lecture title")
        pt = fit_size(lines, w, 2.1, max_pt=30, min_pt=22, where="title slide")
        write_paras(tf, lines, pt, t["title_text"], space_after=0)
        for p in tf.paragraphs:
            for r in p.runs:
                r.font.bold = True
        _, tf = textbox(s, x, 5.05, w, 0.4)
        write_paras(tf, [self.meta["date"]], 16, t["title_muted"], space_after=0)
        _, tf = textbox(s, x, 5.9, w, 0.7)
        write_paras(tf, [self.course["instructor"], self.course["instructor_title"]], 12, t["title_muted"], space_after=2)
        if img:
            picture_contain(s, self.path(img["path"]), W / 2 + 0.5, 0.9, W / 2 - 1.0, 5.4)
            self.caption(s, W / 2 + 0.5, 6.5, W / 2 - 1.0, img["caption"])
        s.notes_slide.notes_text_frame.text = (
            f"{title}\nLecture {self.meta['n']}. {self.course['course']}. {self.meta['date']}."
            + (f"\n\nFigure: {img['caption']}" if img else "")
        )

    def content_slide(self, sd):
        s = self.new_slide()
        where = f"slide {self.num} ({sd.get('title','')!r})"
        for k in ("title", "body", "cite", "refs"):
            if not sd.get(k):
                raise SpecError(f"{where}: missing '{k}'")
        self.title(s, sd["title"])
        lay = sd.get("layout", "text")
        avail = BODY_BOTTOM - BODY_Y
        if lay == "text":
            self.body(s, sd["body"], M, BODY_Y, W - 2 * M, avail, where)
        elif lay in ("figure-right", "figure-left"):
            fig_w = sd.get("figure_width", 5.6)
            text_w = W - 2 * M - fig_w - 0.4
            fx = W - M - fig_w if lay == "figure-right" else M
            tx = M if lay == "figure-right" else M + fig_w + 0.4
            self.body(s, sd["body"], tx, BODY_Y, text_w, avail, where)
            ph = picture_contain(s, self.path(sd["figure"]["path"]), fx, BODY_Y + 0.05, fig_w, avail - 0.75)
            self.caption(s, fx, BODY_Y + 0.15 + ph, fig_w, sd["figure"]["caption"])
        elif lay == "figure-below":
            top_h = sd.get("text_height", 1.9)
            self.body(s, sd["body"], M, BODY_Y, W - 2 * M, top_h, where)
            fy = BODY_Y + top_h + 0.15
            ph = picture_contain(s, self.path(sd["figure"]["path"]), M, fy, W - 2 * M, BODY_BOTTOM - fy - 0.5)
            self.caption(s, M, fy + ph + 0.08, W - 2 * M, sd["figure"]["caption"])
        elif lay == "two-figures":
            top_h = sd.get("text_height", 1.8)
            self.body(s, sd["body"], M, BODY_Y, W - 2 * M, top_h, where)
            figs = sd["figures"]
            if len(figs) != 2:
                raise SpecError(f"{where}: two-figures needs exactly 2 figures")
            fy = BODY_Y + top_h + 0.15
            gap, fw = 0.4, (W - 2 * M - 0.4) / 2
            for i, f in enumerate(figs):
                fx = M + i * (fw + gap)
                ph = picture_contain(s, self.path(f["path"]), fx, fy, fw, BODY_BOTTOM - fy - 0.55)
                self.caption(s, fx, fy + ph + 0.08, fw, f["caption"])
        elif lay == "table":
            tbl = sd["table"]
            rows = len(tbl["rows"]) + 1
            top_h = sd.get("text_height", 2.0)
            self.body(s, sd["body"], M, BODY_Y, W - 2 * M, top_h, where)
            ty = BODY_Y + top_h + 0.15
            row_h = min(0.6, (BODY_BOTTOM - ty) / rows)
            shape = s.shapes.add_table(rows, len(tbl["header"]), Inches(M), Inches(ty),
                                       Inches(W - 2 * M), Inches(row_h * rows))
            table = shape.table
            for r_i, row in enumerate([tbl["header"]] + tbl["rows"]):
                for c_i, val in enumerate(row):
                    cell = table.cell(r_i, c_i)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = rgb(self.t["tint"] if r_i == 0 else self.t["bg"])
                    cell.margin_left = cell.margin_right = Inches(0.08)
                    tf = cell.text_frame
                    tf.word_wrap = True
                    tf.paragraphs[0].text = ""
                    write_paras(tf, [f"**{val}**" if r_i == 0 else val], 13, self.t["text"],
                                bold_color=self.t["heading"], space_after=0)
        else:
            raise SpecError(f"{where}: unknown layout {lay!r}")
        self.footer(s, sd["cite"])
        self.notes(s, sd)

    def takeaways_slide(self, tk):
        s = self.new_slide()
        self.title(s, "Key takeaways")
        items = tk["items"]
        if not 5 <= len(items) <= 6:
            raise SpecError("key takeaways: give 5 or 6 items")
        cols, gap = 2, 0.3
        rows = math.ceil(len(items) / cols)
        cw = (W - 2 * M - gap) / cols
        ch = (BODY_BOTTOM - BODY_Y - gap * (rows - 1)) / rows
        for i, it in enumerate(items):
            cx = M + (i % cols) * (cw + gap)
            cy = BODY_Y + (i // cols) * (ch + gap)
            rect(s, cx, cy, cw, ch, self.t["tint"])
            paras = [f"**{it['lead']}** {it['text']}"]
            pt = fit_size(paras, cw - 0.4, ch - 0.3, max_pt=16, min_pt=12, where=f"takeaway {i+1}")
            _, tf = textbox(s, cx + 0.2, cy + 0.15, cw - 0.4, ch - 0.3)
            write_paras(tf, paras, pt, self.t["text"], bold_color=self.t["heading"])
        self.footer(s, tk["cite"])
        s.notes_slide.notes_text_frame.text = "\n\n".join(tk.get("refs", []))

    def build(self, out):
        want = self.spec.get("content_slides", DEFAULT_CONTENT_SLIDES)
        got = len(self.spec["slides"])
        if got != want:
            raise SpecError(f"spec has {got} content slides; this lecture needs exactly {want}")
        self.title_slide()
        for sd in self.spec["slides"]:
            self.content_slide(sd)
        self.takeaways_slide(self.spec["takeaways"])
        self.prs.core_properties.title = self.meta["title"]
        self.prs.core_properties.author = self.course["instructor"]
        self.prs.save(out)


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    spec_path = Path(sys.argv[1]).resolve()
    spec = json.loads(spec_path.read_text())
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else spec_path.parent / f"Neuroethology_Lecture{spec['lecture']}_FA2026.pptx"
    try:
        Deck(spec, spec_path.parent).build(out)
    except SpecError as e:
        sys.exit(f"BUILD FAILED: {e}")
    print(f"wrote {out}")
    rc = subprocess.call([sys.executable, str(ROOT / "tools/check_lecture.py"), str(out),
                          "--lecture", str(spec["lecture"]),
                          "--content-slides", str(spec.get("content_slides", DEFAULT_CONTENT_SLIDES))])
    sys.exit(rc)


if __name__ == "__main__":
    main()
