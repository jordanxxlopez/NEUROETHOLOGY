#!/usr/bin/env python3
"""Add scaffold slides and third-paragraph fixes to a finished NEUR 411 deck.

    python tools/scaffold_deck.py deck.pptx spec.json out_SCAFFOLDED.pptx

The spec follows the approved scaffolded Lectures 1-10 (see scaffolding/RULES.md):

{
  "third_paragraph": {"4": "replacement for the third body paragraph of slide 4",
                      "6": {"keep": 2, "add": "new teaching sentence(s) after the first 2 sentences"}, ...},
  "inserts": [
    {"after": 4,                       # original slide number the new slide follows
     "title": "Length and width pull prey value in opposite ways",
     "paragraphs": ["...", "...", "..."],      # **bold** marks a key term, _italic_ a species name
     "citation": "Ewert (1970)",               # footer text
     "pictures": [0],                          # optional: which of the parent's figures to keep
     "notes": ["intro sentence", ["Heading.", "sub-bullet", "sub-bullet"], ...],
     "refs_from": ["Clark 2011", ...]}         # full references looked up in the deck's own notes
                                               # (or give "references": ["full reference", ...])
  ]
}

An inserted slide is a copy of the slide it follows (same figure, caption and layout) with a new
title, paragraphs, footer citation, slide number ("04a") and bulleted speaker notes. Every
formatting choice is copied from the parent slide, so each deck keeps its own look. The slide it
follows keeps its number; "third_paragraph" replaces that slide's third body paragraph (where the
closing caveat sat) in the slide and in its speaker-note transcript.
"""
import copy
import json
import re
import sys
from pathlib import Path

from lxml import etree
from PIL import ImageFont
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Emu

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
A, P, R = ("{%s}" % NS[k] for k in ("a", "p", "r"))

FONT = "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"
EMU_PT = 12700
WRAP_FACTOR = 0.94      # wrap a little early so a box never ends up shorter than its text
LINE_FACTOR = 1.15      # rendered line height of Arial is about 1.15 em before line spacing
FOOTER_Y = 6000000      # anything below this line is footer (citation, slide number, rule)
PLAIN = re.compile(r"\*\*|(?<![A-Za-z0-9])_|_(?![A-Za-z0-9])")


def _font(pt):
    return ImageFont.truetype(FONT, size=max(1, int(pt * 10)))   # 10x for sub-point precision


def wrap_lines(text, width_emu, pt):
    font, width = _font(pt), width_emu / EMU_PT * 10
    lines, cur = 1, ""
    for word in text.split():
        trial = (cur + " " + word).strip()
        if font.getlength(trial) <= width:
            cur = trial
        else:
            lines, cur = lines + 1, word
    return lines


# ------------------------------------------------------------------ reading a slide
def classify(slide):
    """Split a content slide into title, body boxes, captions, citation, number and pictures."""
    out = {"title": None, "body": [], "caption": [], "citation": None, "number": None,
           "figs": [], "footer_pics": []}
    pics = [s for s in slide.shapes if s.shape_type == 13]
    out["figs"] = sorted([p for p in pics if p.top < FOOTER_Y], key=lambda s: (s.left, s.top))
    out["footer_pics"] = [p for p in pics if p.top >= FOOTER_Y]
    texts = [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
    for sh in sorted(texts, key=lambda s: s.top):
        t = sh.text_frame.text.strip()
        if sh.top < 900000 and out["title"] is None:
            out["title"] = sh
        elif re.fullmatch(r"\d{1,2}a?", t) and sh.top > FOOTER_Y:
            out["number"] = sh
        elif sh.top > FOOTER_Y and re.search(r"\bFig\.", t):      # a figure caption set low on the slide
            out["caption"].append(sh)
        elif sh.top > FOOTER_Y:
            out["citation"] = sh
        elif sh.name == "Figure caption" or _under_picture(sh, out["figs"]) or \
                (not out["figs"] and sh.top > 5200000 and sh.height < 500000):
            out["caption"].append(sh)
        else:
            out["body"].append(sh)
    out["body"].sort(key=lambda s: (s.left > 6000000, s.top))
    return out


def _under_picture(sh, figs):
    """True when `sh` sits below a figure in the figure's column (so it is that figure's caption)."""
    return any(sh.top >= p.top + p.height - 150000 and p.left - 1600000 <= sh.left <= p.left + p.width
               for p in figs)


def body_paragraphs(slide):
    """All body paragraphs of a slide, in reading order (title and captions excluded)."""
    c = classify(slide)
    return [(sh, p) for sh in c["body"] for p in sh.text_frame.paragraphs if p.text.strip()]


# ------------------------------------------------------------------ writing text
def parse_inline(text):
    """'**bold** and _italic_' -> [(chunk, bold, italic)]."""
    out = []
    for part in re.split(r"(\*\*.+?\*\*|(?<![A-Za-z0-9])_[^_]+?_(?![A-Za-z0-9]))", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            out.append((part[2:-2], True, False))
        elif part.startswith("_") and part.endswith("_") and len(part) > 2:
            out.append((part[1:-1], False, True))
        else:
            out.append((part, False, False))
    return out


def run_templates(c):
    """(normal run, bold run) XML elements to clone, taken from the parent's own body text."""
    normal = bold = None
    for sh in c["body"]:
        for r in sh._element.iter(A + "r"):
            rpr = r.find(A + "rPr")
            is_bold = rpr is not None and rpr.get("b") == "1"
            if is_bold and bold is None:
                bold = r
            if not is_bold and normal is None:
                normal = r
    if bold is None and c["title"] is not None:       # derive from the title (darker theme shade)
        bold = copy.deepcopy(c["title"]._element.find(".//" + A + "r"))
        sz = normal.find(A + "rPr").get("sz")
        if sz:
            bold.find(A + "rPr").set("sz", sz)
    return normal, bold


def make_run(normal, text, bold, italic, bold_template):
    r = copy.deepcopy(bold_template if (bold and bold_template is not None) else normal)
    rpr = r.find(A + "rPr")
    rpr.set("b", "1" if bold else "0")
    if italic:
        rpr.set("i", "1")
    elif rpr.get("i") is not None:
        rpr.set("i", "0")
    r.find(A + "t").text = text
    return r


def fill_paragraph(p_el, text, normal, bold):
    """Rewrite paragraph p_el in place: keep its properties, replace the runs."""
    for child in list(p_el):
        if child.tag in (A + "r", A + "br", A + "fld"):
            p_el.remove(child)
    end = p_el.find(A + "endParaRPr")
    for chunk, b, i in parse_inline(text):
        run = make_run(normal, chunk, b, i, bold)
        if end is not None:
            end.addprevious(run)
        else:
            p_el.append(run)


def set_plain(shape, text):
    """Replace a shape's text with one run, keeping the first run's formatting."""
    p = shape.text_frame.paragraphs[0]._p
    runs = p.findall(A + "r")
    for extra in runs[1:]:
        p.remove(extra)
    runs[0].find(A + "t").text = text
    for extra in shape.text_frame.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def body_pt(box):
    for r in box._element.iter(A + "rPr"):
        if r.get("sz"):
            return int(r.get("sz")) / 100
    return 14.25


def line_spacing(box):
    sp = box._element.find(".//" + A + "lnSpc/" + A + "spcPct")
    return int(sp.get("val")) / 100000 if sp is not None else 1.08


def stack_boxes(boxes, texts):
    """Google-export decks: one box per paragraph, height = wrapped lines, 12 pt gaps."""
    pt, ls = body_pt(boxes[0]), line_spacing(boxes[0])
    top = boxes[0].top
    for box, t in zip(boxes, texts):
        lines = wrap_lines(PLAIN.sub("", t), int(box.width * WRAP_FACTOR), pt)
        h = int(round(lines * pt * LINE_FACTOR * ls * EMU_PT))
        box.top, box.height = Emu(top), Emu(h)
        top += h + 152400


def fit_single_box(box, texts):
    """Single-box decks: shrink the font half a point at a time until the text fits the box."""
    pt, ls = body_pt(box), line_spacing(box)
    sp = box._element.find(".//" + A + "spcAft/" + A + "spcPts")
    gap = int(sp.get("val")) / 100 if sp is not None else 0
    plain = [PLAIN.sub("", t) for t in texts]
    cur = pt
    while cur > 11:
        need = sum(wrap_lines(t, int(box.width * WRAP_FACTOR), cur) * cur * LINE_FACTOR * ls + gap for t in plain) * EMU_PT
        if need <= box.height:
            break
        cur -= 0.5
    if cur != pt:
        for r in box._element.iter(A + "rPr"):
            if r.get("sz"):
                r.set("sz", str(int(cur * 100)))


def write_body(slide, c, texts, normal, bold):
    """Put `texts` into the body, one paragraph each, in the parent's own box structure."""
    boxes = c["body"]
    if len(boxes) == 1:                                   # one box holds every paragraph
        box = boxes[0]
        proto = [p for p in box.text_frame.paragraphs if p.text.strip()][0]._p
        tx = box.text_frame._txBody
        for p in tx.findall(A + "p"):
            tx.remove(p)
        for t in texts:
            np_ = copy.deepcopy(proto)
            fill_paragraph(np_, t, normal, bold)
            tx.append(np_)
        fit_single_box(box, texts)
        return
    for k, t in enumerate(texts):                         # one box per paragraph
        if k >= len(boxes):                               # more paragraphs than boxes: clone the last
            new_el = copy.deepcopy(boxes[-1]._element)
            boxes[-1]._element.addnext(new_el)
            boxes.append([s for s in slide.shapes if s._element is new_el][0])
        fill_paragraph(boxes[k].text_frame.paragraphs[0]._p, t, normal, bold)
    for extra in boxes[len(texts):]:
        extra._element.getparent().remove(extra._element)
    del boxes[len(texts):]
    stack_boxes(boxes, texts)


# ------------------------------------------------------------------ notes
def notes_style(parent):
    """'literal' when the parent's notes use typed bullets ("• "), else real bullet paragraphs."""
    if parent.has_notes_slide:
        for p in parent.notes_slide.notes_text_frame.paragraphs:
            if p.text.lstrip().startswith("• "):
                return "literal"
    return "real"


def parent_figure_notes(parent):
    """The parent's 'Figure:' and 'Figure source:' lines, reused because the figure is reused."""
    if not parent.has_notes_slide:
        return []
    text = parent.notes_slide.notes_text_frame.text
    return [l for l in text.splitlines() if l.startswith(("Figure:", "Figure source:", "Image:"))]


def write_notes(slide, style, items, references, figure_lines):
    body = slide.notes_slide.notes_text_frame._txBody
    for p in body.findall(A + "p"):
        body.remove(p)

    def para(text, lvl, bullet):
        p = etree.SubElement(body, A + "p")
        shown = text
        if style == "real":
            mar = (342900, 685800)[lvl] if bullet else 0
            ppr = etree.SubElement(p, A + "pPr", indent="-228600" if bullet else "0", lvl=str(lvl),
                                   marL=str(mar), marR="0", rtl="0", algn="l")
            for tag in ("spcBef", "spcAft"):
                etree.SubElement(etree.SubElement(ppr, A + tag), A + "spcPts", val="0")
            if bullet:
                etree.SubElement(ppr, A + "buChar", char="•" if lvl == 0 else "–")
            else:
                etree.SubElement(ppr, A + "buNone")
        elif bullet:                                      # typed bullets, as the repo builder writes them
            shown = ("• " if lvl == 0 else "– ") + text
        r = etree.SubElement(p, A + "r")
        rpr = etree.SubElement(r, A + "rPr", lang="en-US", sz="1200")
        if style == "real":
            etree.SubElement(etree.SubElement(rpr, A + "solidFill"), A + "schemeClr", val="dk1")
            for tag in ("latin", "ea", "cs", "sym"):
                etree.SubElement(rpr, A + tag, typeface="Calibri")
        etree.SubElement(r, A + "t").text = shown

    for item in items:
        if isinstance(item, str):
            para(item, 0, True)
        else:
            para(item[0], 0, True)
            for sub in item[1:]:
                para(sub, 1, True)
    para("", 0, False)
    para("References:", 0, False)
    for ref in references:
        para(ref, 0, False)
    if figure_lines:
        para("", 0, False)
        for line in figure_lines:
            para(line, 0, False)


# ------------------------------------------------------------------ slides
def clone_slide(prs, parent):
    """Copy `parent` (shapes, background, pictures) into a new slide at the end of the deck."""
    new = prs.slides.add_slide(parent.slide_layout)
    for sh in list(new.shapes):
        sh._element.getparent().remove(sh._element)
    bg = parent._element.find(".//" + P + "bg")
    if bg is not None:
        new._element.find(P + "cSld").insert(0, copy.deepcopy(bg))
    tree = new.shapes._spTree
    for el in parent.shapes._spTree:
        if el.tag in (P + "nvGrpSpPr", P + "grpSpPr"):
            continue
        tree.append(copy.deepcopy(el))
    for blip in tree.iter(A + "blip"):
        rid = blip.get(R + "embed")
        if rid:
            blip.set(R + "embed", new.part.relate_to(parent.part.related_part(rid), RT.IMAGE))
    return new


def move_after(prs, index):
    """Move the slide just added at the end of the deck so it follows slide `index`."""
    lst = prs.slides._sldIdLst
    el = list(lst)[-1]
    lst.remove(el)
    lst.insert(index + 1, el)


def add_scaffold(prs, parent_idx, spec, number):
    parent = prs.slides[parent_idx]
    style = notes_style(parent)
    normal, bold = run_templates(classify(parent))
    new = clone_slide(prs, parent)
    c = classify(new)
    for pic in c["footer_pics"]:                          # the approved scaffolded slides omit it
        pic._element.getparent().remove(pic._element)
    if "pictures" in spec:
        for k, pic in enumerate(c["figs"]):
            if k not in spec["pictures"]:
                for cap in [s for s in c["caption"] if _under_picture(s, [pic])]:
                    cap._element.getparent().remove(cap._element)
                pic._element.getparent().remove(pic._element)
    set_plain(c["title"], spec["title"])
    write_body(new, c, spec["paragraphs"], normal, bold)
    set_plain(c["citation"], spec["citation"])
    need = int(_font(8.25).getlength(spec["citation"]) / 10 * EMU_PT * 1.15)
    c["citation"].width = Emu(max(c["citation"].width, need, 700000))
    old = c["number"].text_frame.text.strip()
    label = f"{number:02d}a" if len(old) == 2 else f"{number}a"
    right = c["number"].left + c["number"].width
    set_plain(c["number"], label)
    c["number"].width = Emu(c["number"].width + 62417)
    c["number"].left = Emu(right - c["number"].width)
    if new.notes_slide.notes_text_frame is None:         # notes master without placeholders: reuse the parent's
        tree = new.notes_slide.shapes._spTree
        for el in parent.notes_slide.shapes._spTree:
            if el.tag not in (P + "nvGrpSpPr", P + "grpSpPr"):
                tree.append(copy.deepcopy(el))
    write_notes(new, style, spec["notes"], spec["references"], [])
    move_after(prs, parent_idx)
    return new


SENTENCE = re.compile(r"(?<=[.!?])\s+(?=[A-Z])")


def resolve_third(old, value):
    """A fix is either the full replacement text or {"keep": n, "add": text}: the first n sentences
    of the existing third paragraph followed by new teaching sentences."""
    if isinstance(value, str):
        return value
    sentences = SENTENCE.split(old.strip())
    return " ".join(sentences[:value["keep"]] + [value["add"]])


def fix_third_paragraph(slide, value):
    """Replace the slide's third body paragraph (where the caveat sat) in the slide and its notes.

    Only that paragraph is rewritten; the other paragraphs keep their runs, bold terms included."""
    c = classify(slide)
    paras = body_paragraphs(slide)
    if len(paras) < 3:
        raise SystemExit(f"slide has {len(paras)} body paragraphs; cannot fix the third")
    normal, bold = run_templates(c)
    box, para = paras[2]
    old = para.text.strip()
    new_text = resolve_third(old, value)
    fill_paragraph(para._p, new_text, normal, bold)
    texts = [PLAIN.sub("", p.text) for _, p in body_paragraphs(slide)]
    if len(c["body"]) == 1:
        fit_single_box(c["body"][0], texts)
    else:
        stack_boxes(c["body"], texts)
    if slide.has_notes_slide:
        for t in slide.notes_slide.notes_text_frame._txBody.iter(A + "t"):
            if t.text and old[:50] in t.text:
                t.text = t.text.replace(old, PLAIN.sub("", new_text))
    return old


def reference_index(prs):
    """Every full reference line found after 'References:' in the deck's own speaker notes."""
    refs = []
    for sl in prs.slides:
        if not sl.has_notes_slide:
            continue
        lines = sl.notes_slide.notes_text_frame.text.splitlines()
        if "References:" in lines:
            for line in lines[lines.index("References:") + 1:]:
                if not line.strip():
                    break
                if line not in refs:
                    refs.append(line)
    return refs


def resolve_refs(index, keys):
    """'Clark 2011' -> the deck's full reference for Clark ... (2011)."""
    out = []
    for key in keys:
        name, year = key.rsplit(" ", 1)
        hits = [r for r in index if (r.startswith(name + " ") or r.startswith(name + ",")) and f"({year})" in r[:900]]
        if not hits:                                    # "MS Brainard; EI Knudsen (1993)": first author, initials first
            hits = [r for r in index
                    if re.search(r"\b" + re.escape(name) + r"\b", re.split(r"[;(]", r, 1)[0]) and f"({year})" in r[:900]]
        if not hits:
            raise SystemExit(f"no reference in the deck's notes for {key!r}")
        out.append(hits[0])
    return out


def main():
    deck, spec_path, out = sys.argv[1:4]
    spec = json.loads(Path(spec_path).read_text())
    prs = Presentation(deck)
    n_orig = len(prs.slides)
    for num, text in spec.get("third_paragraph", {}).items():
        fix_third_paragraph(prs.slides[int(num) - 1], text)
    index = reference_index(prs)
    for ins in spec.get("inserts", []):
        ins["references"] = resolve_refs(index, ins["refs_from"]) if "refs_from" in ins else ins["references"]
    for ins in sorted(spec.get("inserts", []), key=lambda i: -i["after"]):   # back to front
        add_scaffold(prs, ins["after"] - 1, ins, ins["after"])
    prs.save(out)
    print(f"{deck}: {n_orig} -> {len(prs.slides)} slides, "
          f"{len(spec.get('third_paragraph', {}))} third paragraphs fixed -> {out}")


if __name__ == "__main__":
    main()
