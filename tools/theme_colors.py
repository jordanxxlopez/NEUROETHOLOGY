#!/usr/bin/env python3
"""One color per lecture: titles and bold text are a darker shade of the title-slide color;
all other text is black.

The lecture's identity color is the title-slide panel (or, when the panel is pale, the
colored title-slide text). Slide titles, bold terms and takeaway lead phrases use a darker
shade of that color; tints and hairlines are pale versions of it; body text, captions and
footers are black (000000). build_lecture.py applies this to every palette, and
check_lecture.py fails any deck that breaks it.

Fix an existing deck (e.g. one recolored in Google Slides):
    python tools/theme_colors.py deck.pptx -o fixed.pptx
"""
import argparse
import colorsys
from collections import Counter

BLACK = "000000"
TITLE_PT = 24          # bold runs at or above this size on content slides are slide titles
HUE_TOLERANCE = 0.06   # fraction of the color wheel (~22 degrees)


def _hls(hex_):
    r, g, b = (int(hex_[i:i + 2], 16) / 255 for i in (0, 2, 4))
    return colorsys.rgb_to_hls(r, g, b)


def _hex(h, l, s):
    return "".join(f"{round(v * 255):02X}" for v in colorsys.hls_to_rgb(h, l, s))


def identity_color(title_bg, title_text, title_muted=None):
    """The color that defines the lecture: a strong title panel, else its colored text."""
    if _hls(title_bg)[1] < 0.6:
        return title_bg
    cands = [c for c in (title_text, title_muted) if c]
    return max(cands, key=lambda c: (_hls(c)[2] > 0.15, _hls(c)[2]))


def darker(hex_):
    h, l, s = _hls(hex_)
    return _hex(h, min(l * 0.8, 0.30), min(s, 0.7))


def pale(hex_, lightness, max_sat):
    h, _, s = _hls(hex_)
    return _hex(h, lightness, min(s, max_sat))


def aligned_palette(t):
    """Palette with heading/bold, tint and rule derived from the title-slide color; text black."""
    ident = identity_color(t["title_bg"], t["title_text"], t.get("title_muted"))
    out = dict(t)
    out.update(heading=darker(ident), text=BLACK, muted=BLACK,
               tint=pale(ident, 0.96, 0.35), rule=pale(ident, 0.84, 0.25))
    return out


# ------------------------------------------------------------- finished decks
def _run_color(run):
    try:
        return str(run.font.color.rgb)
    except Exception:
        return None


def _frames(shapes):
    for sh in shapes:
        if sh.shape_type == 6:  # group
            yield from _frames(sh.shapes)
        elif sh.has_text_frame:
            yield sh.text_frame
        elif getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for cell in row.cells:
                    yield cell.text_frame


def _solid_fill(sh):
    try:
        if sh.fill.type == 1:
            return str(sh.fill.fore_color.rgb)
    except Exception:
        pass
    return None


def deck_identity(prs):
    """Identity color of a finished deck, read from its title slide."""
    s1 = prs.slides[0]
    panels = [(sh.width * sh.height, c) for sh in s1.shapes if (c := _solid_fill(sh))]
    panel = max(panels)[1] if panels else None
    if panel is None:
        try:
            panel = str(s1.background.fill.fore_color.rgb)
        except Exception:
            panel = "FFFFFF"
    text = Counter(c for tf in _frames(s1.shapes) for p in tf.paragraphs for r in p.runs
                   if r.text.strip() and (c := _run_color(r)))
    colored = [c for c, _ in text.most_common() if _hls(c)[2] > 0.15 and _hls(c)[1] < 0.6]
    if _hls(panel)[1] < 0.6:
        return panel
    return colored[0] if colored else (text.most_common(1)[0][0] if text else "333333")


def align_deck(prs):
    """Recolor content slides in place; the title slide is left as designed."""
    ident = deck_identity(prs)
    heading, tint, rule = darker(ident), pale(ident, 0.96, 0.35), pale(ident, 0.84, 0.25)
    for slide in list(prs.slides)[1:]:
        for tf in _frames(slide.shapes):
            for p in tf.paragraphs:
                for r in p.runs:
                    if r.text.strip():
                        r.font.color.rgb = _rgb(heading if r.font.bold else BLACK)
        for sh in slide.shapes:
            c = _solid_fill(sh)
            if c and c != "FFFFFF" and sh.shape_type != 13:
                light = _hls(c)[1]
                if light >= 0.9:
                    sh.fill.fore_color.rgb = _rgb(tint)
                elif light >= 0.75:
                    sh.fill.fore_color.rgb = _rgb(rule)
    return ident, heading


def alignment_problems(prs):
    """Rule violations in a finished deck (titles/bold off-color, body text not black)."""
    ident = deck_identity(prs)
    ih, il, is_ = _hls(ident)
    bad_bold, bad_text = Counter(), Counter()
    for i, slide in enumerate(list(prs.slides)[1:], start=2):
        for tf in _frames(slide.shapes):
            for p in tf.paragraphs:
                for r in p.runs:
                    c = _run_color(r)
                    if not r.text.strip() or c is None:
                        continue
                    if r.font.bold:
                        h, l, s = _hls(c)
                        dh = min(abs(h - ih), 1 - abs(h - ih))
                        same_hue = (s < 0.12 and is_ < 0.12) or (dh <= HUE_TOLERANCE and is_ >= 0.12)
                        if not same_hue or l > max(il, 0.35) + 0.02:
                            bad_bold[i] += 1
                    elif c != BLACK:
                        bad_text[i] += 1
    out = []
    if bad_bold:
        out.append(f"titles/bold not a darker shade of the title-slide color #{ident} on slides "
                   f"{sorted(bad_bold)[:8]}")
    if bad_text:
        out.append(f"non-bold text is not black on slides {sorted(bad_text)[:8]}")
    return out


def _rgb(hex_):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(hex_)


def main():
    from pptx import Presentation
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("deck")
    ap.add_argument("-o", "--out", required=True)
    a = ap.parse_args()
    prs = Presentation(a.deck)
    ident, heading = align_deck(prs)
    prs.save(a.out)
    print(f"title-slide color #{ident} -> titles/bold #{heading}; other text black; wrote {a.out}")


if __name__ == "__main__":
    main()
